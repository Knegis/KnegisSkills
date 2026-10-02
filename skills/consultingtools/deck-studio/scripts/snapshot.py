"""Capture a slide - from a .pptx or a standalone .html file - as a PNG and file it in the
asset bank, so an agent can later place it inside another slide via <img src="..."> (see
"Images" in PIPELINE.md). This is the slide-in-slide library: a capture of one slide used as
a picture inside a different slide.

Usage:
  python snapshot.py <deck.pptx> <N> [--out file.png] [--desc "..."] [--tags a,b] [--width 1920]
  python snapshot.py <slide.html> [--out file.png] [--desc "..."] [--tags a,b]
  python snapshot.py <deck.pptx> --all [--desc-prefix "..."] [--tags a,b] [--width 1920]

<N> is the 1-based slide number. Without --out the PNG is filed in the asset bank's
screenshot/ library (../asset_bank beside the skill) under a fresh aNNNN id, with a MANIFEST.json record
alongside every other queryable asset (find_asset.py searches them).
With --out the PNG is written to that path instead and no manifest record is created.

.pptx capture uses PowerPoint COM (Slides(N).Export) to export just that slide when available,
falling back to rendering the whole deck via render_slides.py's Windows/macOS/Linux dispatch
and picking slide N - a deck you have open is never touched, and only a PowerPoint instance
this script itself started is ever quit. .html capture reuses render_html.py's browser-launch
helper at 2x device scale (Edge, then Chrome, then bundled Chromium).

--all captures every slide of the deck in one pass, one manifest record per slide.

Filed captures are deduped by SHA-256: capturing an identical image a second time prints the
existing id and does not add a second manifest record.
"""
import hashlib
import html
import json
import platform
import re
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import _cli  # noqa: E402 - shared --help / arg-count handling
import render_slides  # noqa: E402 - reuse its multi-platform render dispatch as the fallback path
import render_html as _render_html  # noqa: E402 - reuse its browser-launch helpers, do not shell out

ASSET_BANK = HERE.parent.parent / "asset_bank"
SCREENSHOT_DIR = "screenshot"
DEFAULT_WIDTH = 1920
HTML_DEVICE_SCALE = 2   # matches render_html.py's fixed device_scale_factor


# ============================================================================
# manifest helpers
# ============================================================================

def _manifest_path():
    return ASSET_BANK / "MANIFEST.json"


def load_manifest():
    # No asset bank yet (none ships with the skill): start an empty one on first capture.
    if not _manifest_path().exists():
        return {"assets": []}
    return json.loads(_manifest_path().read_text(encoding="utf-8"))


def save_manifest(manifest):
    # Match the existing file's formatting exactly (indent=1, raw UTF-8, no trailing
    # newline) so a diff of MANIFEST.json shows only the appended record.
    _manifest_path().write_text(
        json.dumps(manifest, indent=1, ensure_ascii=False), encoding="utf-8"
    )


def next_id(assets):
    """aNNNN, one past the highest existing id (a0526 -> a0527)."""
    highest = 0
    for a in assets:
        m = re.match(r"^a(\d+)$", a.get("id", ""))
        if m:
            highest = max(highest, int(m.group(1)))
    return f"a{highest + 1:04d}"


def find_by_sha(assets, sha256):
    for a in assets:
        if a.get("sha256") == sha256:
            return a
    return None


def slugify(text, maxlen=50):
    s = re.sub(r"[^a-z0-9]+", "_", (text or "").lower()).strip("_")
    s = re.sub(r"_+", "_", s)
    return s[:maxlen].strip("_") or "capture"


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def png_dims(path):
    from PIL import Image
    with Image.open(path) as im:
        return list(im.size)


def default_description(title, deck_stem, slide_no):
    if slide_no:
        return f"Slide capture: {title} ({deck_stem}, slide {slide_no})"
    return f"Slide capture: {title} ({deck_stem})"


def build_record(id_, rel_file, deck_stem, source_str, slide_no, title, desc, tags, png_path, taken):
    """One MANIFEST.json record for a filed capture. png_path must already exist at rel_file's
    final location (or an identical temp copy) - dims/bytes/sha are read from it directly."""
    merged_tags = ["snapshot", "slide capture", deck_stem]
    for t in (tags or []):
        if t not in merged_tags:
            merged_tags.append(t)
    w, h = png_dims(png_path)
    return {
        "id": id_,
        "file": rel_file,
        "category": "screenshot",
        "ext": "png",
        "bytes": Path(png_path).stat().st_size,
        "dims": [w, h],
        "alpha": False,
        "sources": [f"{source_str} slide {slide_no}" if slide_no else source_str],
        "source_slides": [],
        "slide_titles": [title],
        "pptx_names": [],
        "descr": [],
        "description": desc or default_description(title, deck_stem, slide_no),
        "tags": merged_tags,
        "sha256": sha256_file(png_path),
        "prev_file": "",
        "snapshot": {"source": source_str, "slide": slide_no, "taken": taken},
    }


def file_in_bank(png_path, source_str, slide_no, title, desc, tags, deck_stem):
    """Dedupe-then-file a rendered PNG into the asset bank. Returns (record, created) -
    created is False when an identical image already existed (record is the existing one)."""
    (ASSET_BANK / SCREENSHOT_DIR).mkdir(parents=True, exist_ok=True)
    sha = sha256_file(png_path)
    manifest = load_manifest()
    existing = find_by_sha(manifest["assets"], sha)
    if existing:
        return existing, False

    id_ = next_id(manifest["assets"])
    rel_file = f"{SCREENSHOT_DIR}/{id_}_{slugify(title)}.png"
    dest = ASSET_BANK / rel_file
    shutil.copyfile(png_path, dest)
    taken = datetime.now(timezone.utc).isoformat(timespec="seconds")
    record = build_record(id_, rel_file, deck_stem, source_str, slide_no, title, desc, tags, dest, taken)
    manifest["assets"].append(record)
    save_manifest(manifest)
    return record, True


# ============================================================================
# slide title extraction
# ============================================================================

def _slide_title(slide, n):
    title_shape = None
    try:
        title_shape = slide.shapes.title
    except Exception:
        title_shape = None
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        t = title_shape.text_frame.text.strip()
        if t:
            return t
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            t = shape.text_frame.text.strip()
            if t:
                return t
    return f"slide {n}"


def slide_title_from_pptx(pptx_path, n):
    from pptx import Presentation
    prs = Presentation(str(pptx_path))
    count = len(prs.slides)
    if n < 1 or n > count:
        raise ValueError(f"slide {n} out of range (deck has {count} slides)")
    return _slide_title(prs.slides[n - 1], n)


def _clean_html_fragment(fragment):
    text = re.sub(r"<[^>]+>", " ", fragment)
    text = html.unescape(text)
    return re.sub(r"\s+", " ", text).strip()


def _matching_close_index(html_text, tag, start):
    """start is just past a <tag ...> opening tag's '>'. Depth-aware scan for the
    matching </tag>, so a same-named tag nested inside (e.g. .slide wrapping .title,
    both <div>s) doesn't fool a naive first-</tag> match. Returns None if unmatched."""
    depth = 1
    for m in re.finditer(rf"<(/?){tag}\b[^>]*?(/?)>", html_text[start:], re.IGNORECASE):
        if m.group(1):
            depth -= 1
            if depth == 0:
                return start + m.start()
        elif not m.group(2):
            depth += 1
    return None


def title_from_html(html_text):
    """The text of the first element carrying class="title", else <title>, else "untitled"."""
    for m in re.finditer(r'<([a-zA-Z][a-zA-Z0-9]*)\b([^>]*)>', html_text):
        tag, attrs = m.group(1), m.group(2)
        if attrs.rstrip().endswith("/"):
            continue  # self-closing: no content
        cm = re.search(r'\bclass\s*=\s*"([^"]*)"', attrs)
        if not cm or "title" not in cm.group(1).split():
            continue
        close_start = _matching_close_index(html_text, tag, m.end())
        if close_start is None:
            continue
        text = _clean_html_fragment(html_text[m.end():close_start])
        if text:
            return text
    m = re.search(r"<title[^>]*>(.*?)</title>", html_text, re.DOTALL | re.IGNORECASE)
    if m:
        text = _clean_html_fragment(m.group(1))
        if text:
            return text
    return "untitled"


# ============================================================================
# rendering
# ============================================================================

def _render_pptx_slide_windows_com(pptx_path, n, out_png, width):
    import uuid
    import win32com.client

    tmp = Path(tempfile.gettempdir()) / f"snapshot_render_{uuid.uuid4().hex}.pptx"
    shutil.copyfile(pptx_path, tmp)
    try:
        win32com.client.GetActiveObject("PowerPoint.Application")
        pre_running = True
    except Exception:
        pre_running = False

    ppt = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = ppt.Presentations.Open(str(tmp), -1, 0, 0)
        try:
            count = pres.Slides.Count
            if n < 1 or n > count:
                raise ValueError(f"slide {n} out of range (deck has {count} slides)")
            height = int(width * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
            pres.Slides(n).Export(str(out_png), "PNG", width, height)
        finally:
            pres.Close()
    finally:
        if not pre_running:
            try:
                ppt.Quit()
            except Exception:
                pass
        tmp.unlink(missing_ok=True)


def _render_pptx_slide_via_full_render(pptx_path, n, out_png, width):
    """Fallback: render the whole deck through render_slides.py's own platform dispatch
    (which has its own pywin32 -> comtypes -> .ps1 / LibreOffice fallback chain) to a
    scratch dir, then lift out slide N."""
    work = Path(tempfile.mkdtemp(prefix="snapshot_full_"))
    try:
        system = platform.system()
        if system == "Windows":
            count, how = render_slides._windows(pptx_path, work, width)
        elif system == "Darwin":
            count, how = render_slides._macos(pptx_path, work, width)
        else:
            count, how = render_slides._linux(pptx_path, work, width)
        if not count:
            raise RuntimeError("no slide renderer available on this machine")
        if n < 1 or n > count:
            raise ValueError(f"slide {n} out of range (deck has {count} slides)")
        src = work / f"slide{n:02d}.png"
        if not src.exists():
            raise RuntimeError(f"expected render {src} was not produced")
        shutil.copyfile(src, out_png)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_pptx_slide(pptx_path, n, out_png, width=DEFAULT_WIDTH):
    out_png = Path(out_png)
    out_png.parent.mkdir(parents=True, exist_ok=True)
    if platform.system() == "Windows":
        try:
            import win32com.client  # noqa: F401
        except ImportError:
            pass
        else:
            try:
                _render_pptx_slide_windows_com(pptx_path, n, out_png, width)
                return
            except ValueError:
                raise
            except Exception as e:
                print(f"single-slide PowerPoint COM export failed, falling back to a full "
                      f"render: {e}", file=sys.stderr)
    _render_pptx_slide_via_full_render(pptx_path, n, out_png, width)


def render_pptx_all(pptx_path, out_dir, width=DEFAULT_WIDTH):
    """Render every slide into out_dir (slideNN.png). Returns the slide count."""
    system = platform.system()
    if system == "Windows":
        count, how = render_slides._windows(pptx_path, out_dir, width)
    elif system == "Darwin":
        count, how = render_slides._macos(pptx_path, out_dir, width)
    else:
        count, how = render_slides._linux(pptx_path, out_dir, width)
    if not count:
        raise RuntimeError("no slide renderer available on this machine")
    return count


def render_html_slide(html_path, out_png):
    """2x-scale screenshot of a standalone slide HTML file, reusing render_html.py's
    browser-launch and BOM-stripping helpers rather than shelling out to it."""
    from playwright.sync_api import sync_playwright

    out_png = Path(out_png)
    out_png.parent.mkdir(parents=True, exist_ok=True)
    _render_html._strip_bom(html_path)
    with sync_playwright() as p:
        browser, how = _render_html._launch(p)
        try:
            page = browser.new_page(viewport={"width": 1280, "height": 720},
                                     device_scale_factor=HTML_DEVICE_SCALE)
            page.goto(html_path.as_uri(), wait_until="networkidle",
                       timeout=_render_html.TIMEOUT_MS)
            shim = HERE / "chart_preview.js"
            if shim.exists():
                page.add_script_tag(path=str(shim))
            page.wait_for_timeout(300)
            page.screenshot(path=str(out_png), timeout=_render_html.TIMEOUT_MS)
        finally:
            browser.close()


# ============================================================================
# CLI
# ============================================================================

def _cmd_pptx_slide(pptx_path, n, out, desc, tags, width):
    title = slide_title_from_pptx(pptx_path, n)
    if out:
        out_path = Path(out).resolve()
        render_pptx_slide(pptx_path, n, out_path, width)
        print(f"wrote {out_path}")
        return
    work = Path(tempfile.mkdtemp(prefix="snapshot_"))
    try:
        tmp_png = work / "capture.png"
        render_pptx_slide(pptx_path, n, tmp_png, width)
        record, created = file_in_bank(tmp_png, str(pptx_path), n, title, desc, tags, pptx_path.stem)
        _report(record, created)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _cmd_pptx_all(pptx_path, desc_prefix, tags, width):
    work = Path(tempfile.mkdtemp(prefix="snapshot_all_"))
    try:
        count = render_pptx_all(pptx_path, work, width)
        for n in range(1, count + 1):
            src_png = work / f"slide{n:02d}.png"
            title = slide_title_from_pptx(pptx_path, n)
            desc = f"{desc_prefix}: {title}" if desc_prefix else None
            record, created = file_in_bank(src_png, str(pptx_path), n, title, desc, tags, pptx_path.stem)
            _report(record, created)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _cmd_html(html_path, out, desc, tags):
    text = html_path.read_text(encoding="utf-8", errors="ignore")
    title = title_from_html(text)
    if out:
        out_path = Path(out).resolve()
        render_html_slide(html_path, out_path)
        print(f"wrote {out_path}")
        return
    work = Path(tempfile.mkdtemp(prefix="snapshot_"))
    try:
        tmp_png = work / "capture.png"
        render_html_slide(html_path, tmp_png)
        record, created = file_in_bank(tmp_png, str(html_path), None, title, desc, tags, html_path.stem)
        _report(record, created)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _report(record, created):
    loc = ASSET_BANK / record["file"]
    note = "" if created else " (already in the bank)"
    print(f"{record['id']}  {record.get('description', '')}  {loc}{note}")


def main():
    raw = sys.argv[1:]
    # --out / --desc / --tags / --width / --desc-prefix each take a value that is not
    # itself a flag - drop them before the generic --help / positional-count check, or
    # the value would be miscounted as a positional (see scripts/_cli.py).
    argv_for_cli = list(raw)
    for flag in ("--out", "--desc", "--tags", "--width", "--desc-prefix"):
        while flag in argv_for_cli:
            i = argv_for_cli.index(flag)
            del argv_for_cli[i:i + 2]
    _cli.usage(__doc__, argv_for_cli, 1, 2)

    def flag_value(name):
        if name in raw:
            i = raw.index(name)
            if i + 1 < len(raw):
                return raw[i + 1]
        return None

    out = flag_value("--out")
    desc = flag_value("--desc")
    desc_prefix = flag_value("--desc-prefix")
    tags_raw = flag_value("--tags")
    tags = [t.strip() for t in tags_raw.split(",") if t.strip()] if tags_raw else []
    width_raw = flag_value("--width")
    all_flag = "--all" in raw

    positional = [a for a in raw if not a.startswith("-")]
    src = Path(positional[0])
    if not src.exists():
        print(f"FAILED: not found: {src}")
        sys.exit(1)
    src = src.resolve()

    try:
        width = int(width_raw) if width_raw else DEFAULT_WIDTH
        if src.suffix.lower() == ".pptx":
            if all_flag:
                if out:
                    print("FAILED: --out is not supported with --all")
                    sys.exit(1)
                _cmd_pptx_all(src, desc_prefix, tags, width)
            else:
                if len(positional) < 2:
                    print("FAILED: a .pptx capture needs a slide number (or --all)")
                    sys.exit(1)
                try:
                    n = int(positional[1])
                except ValueError:
                    print(f"FAILED: not a slide number: {positional[1]}")
                    sys.exit(1)
                _cmd_pptx_slide(src, n, out, desc, tags, width)
        elif src.suffix.lower() == ".html":
            if len(positional) > 1:
                print(f"FAILED: unexpected extra argument for an .html capture: {positional[1]}")
                sys.exit(1)
            _cmd_html(src, out, desc, tags)
        else:
            print(f"FAILED: unsupported file type: {src.suffix}")
            sys.exit(1)
    except SystemExit:
        raise
    except Exception as e:
        print(f"FAILED: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
