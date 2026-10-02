"""Add a cover, divider or end slide using the TEMPLATE'S OWN layout and placeholders.

    python scripts/chrome_slide.py <deck.pptx> --layout "<name or alias>" --title "Text" [--subtitle "Text"]
                                   [--body "line 1|line 2"] [--date "2026-09-11"] [--at N] [--new]

Cover, section-divider and end slides carry no argument, so they should look exactly like the template
intends. The HTML pipeline draws its own shapes and would leave the layout's title placeholder empty
underneath; this script instead fills the placeholders with python-pptx, which is what a person does in
PowerPoint. Use it for chrome slides; use the HTML pipeline for content slides.

  <deck.pptx>   an existing deck to append to (default), or a new file with --new (built on the house template)
  --layout      a layout name substring ("Title Slide", "Section Header", "Title - Dark") or alias
                (title | divider | end); run  python scripts/list_layouts.py  to see the names
  --title       goes into the title placeholder
  --subtitle    goes into the first body/subtitle placeholder, if the layout has one
  --body        one or more lines (| separated) for the next body placeholder (e.g. contact block on an end slide)
  --date        fills the date placeholder if present (default: leave the template's field)
  --at N        insert as slide N (1-based) instead of appending
  --new         create <deck.pptx> from the configured template instead of opening it

Every text run is set to the house font. Placeholders you do not fill are removed so no empty prompt boxes
remain.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as C  # noqa: E402
from _cli import usage  # noqa: E402

from pptx import Presentation  # noqa: E402

ALIASES = {"title": ["title -", "title slide", "cover"],
           "divider": ["divider", "section header", "section"],
           "end": ["end slide", "end", "closing", "thank"]}


def _opt(argv, flag, default=None):
    if flag in argv:
        i = argv.index(flag)
        if i + 1 < len(argv):
            v = argv[i + 1]
            del argv[i:i + 2]
            return v
        del argv[i]
    return default


def _norm(s: str) -> str:
    """Lower-case and treat en/em dashes as hyphens, so 'Title - Dark' finds 'Title — Dark'."""
    return s.lower().replace("—", "-").replace("–", "-").replace("  ", " ")


def find_layout(prs, want: str):
    names = [(l, _norm(l.name)) for l in prs.slide_layouts]
    w = _norm(want.strip())
    for pats in ([w] if w not in ALIASES else ALIASES[w],):
        for p in pats:
            hits = [l for l, n in names if p in n]
            if hits:
                pref = [l for l in hits if any(k in _norm(l.name) for k in ("white", "light", "cream"))]
                return (pref or hits)[0]
    raise SystemExit("no layout matches '%s'. Layouts: %s" % (want, ", ".join(l.name for l, _ in names)))


def set_text(ph, text: str):
    tf = ph.text_frame
    lines = [t for t in text.split("|")]
    tf.text = lines[0]
    for extra in lines[1:]:
        p = tf.add_paragraph()
        p.text = extra
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.name = C.FONT_NAME


def _force_house_font(deck: Path):
    """The template's layouts and theme carry fallback typefaces; rewrite them like build_deck does."""
    import tempfile
    import office_io
    import patch_fonts
    with tempfile.TemporaryDirectory() as td:
        unp = Path(td) / "unpacked"
        office_io.unpack(str(deck), str(unp))
        patch_fonts.patch(str(unp), quiet=True) if "quiet" in patch_fonts.patch.__code__.co_varnames else patch_fonts.patch(str(unp))
        out = Path(td) / "packed.pptx"
        office_io.pack(str(unp), str(out))
        shutil.copyfile(out, deck)


def main(argv):
    argv = list(argv)
    usage(__doc__, argv, 1, 99)
    new = "--new" in argv
    if new:
        argv.remove("--new")
    layout = _opt(argv, "--layout", "title")
    title = _opt(argv, "--title", "")
    subtitle = _opt(argv, "--subtitle", None)
    body = _opt(argv, "--body", None)
    date = _opt(argv, "--date", None)
    at = _opt(argv, "--at", None)
    deck = Path(argv[0])
    if new:
        if deck.exists():
            raise SystemExit("%s exists - refusing to overwrite (drop --new to append, or choose a new name)" % deck)
        shutil.copyfile(C.TEMPLATE, deck)
    elif not deck.exists():
        raise SystemExit("not found: %s (use --new to create it from the template)" % deck)

    prs = Presentation(str(deck))
    lay = find_layout(prs, layout)
    slide = prs.slides.add_slide(lay)

    filled = set()
    for ph in list(slide.placeholders):
        t = ph.placeholder_format.type
        tname = str(t).split(".")[-1].split(" ")[0].lower() if t is not None else ""
        if tname in ("title", "center_title") and title:
            set_text(ph, title); filled.add(ph.placeholder_format.idx)
        elif tname in ("subtitle",) and subtitle is not None:
            set_text(ph, subtitle); filled.add(ph.placeholder_format.idx)
        elif tname == "date" and date is not None:
            set_text(ph, date); filled.add(ph.placeholder_format.idx)
    # body placeholders: first gets subtitle if no subtitle placeholder existed, next gets --body
    pending = []
    if subtitle is not None and not any(
            str(p.placeholder_format.type).split(".")[-1].lower().startswith("subtitle") for p in slide.placeholders):
        pending.append(subtitle)
    if body is not None:
        pending.append(body)
    # largest body placeholder first: on an end slide that is the contact block, not the small address box
    bodies = [ph for ph in slide.placeholders if ph.placeholder_format.idx not in filled
              and str(ph.placeholder_format.type).split(".")[-1].split(" ")[0].lower() in ("body", "object")]
    bodies.sort(key=lambda ph: -(ph.width or 0) * (ph.height or 0))
    for ph in bodies:
        if not pending:
            break
        set_text(ph, pending.pop(0)); filled.add(ph.placeholder_format.idx)
    for left in pending:
        print("  ! not placed (layout '%s' has no free body placeholder): %s" % (lay.name, left[:60]))
    # remove empty prompt boxes (keep footer/date/slide-number fields, which the master renders)
    for ph in list(slide.placeholders):
        tname = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0].lower()
        if ph.placeholder_format.idx in filled or tname in ("date", "footer", "slide_number"):
            continue
        ph._element.getparent().remove(ph._element)

    if at is not None:
        n = int(at)
        sldIdLst = prs.slides._sldIdLst
        ids = list(sldIdLst)
        el = ids[-1]
        sldIdLst.remove(el)
        sldIdLst.insert(n - 1, el)

    prs.save(str(deck))
    _force_house_font(deck)
    print("added '%s' slide on layout '%s' -> %s (%d slides)" % (layout, lay.name, deck, len(prs.slides)))


if __name__ == "__main__":
    main(sys.argv[1:])
