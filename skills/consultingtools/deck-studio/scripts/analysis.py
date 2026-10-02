"""Reads for the workroom Analysis tab: the sidecar JSON beside a slide, and the Excel ranges
its figures point back to.

Nothing here ever writes to a workbook. Workbooks are opened read_only=True (openpyxl) and
cached in-process keyed by (path, mtime), once for formulas (data_only=False) and once for
values (data_only=True), so repeated /api/analysis and /api/analysis/sheet calls against the
same file don't reopen it every time. A file that changes on disk (new mtime) evicts its old
cache entry.

See scripts/workroom.py ("Analysis tab") for how this is wired into the server, and PIPELINE.md
/ WORKROOM.md for the review loop this supports. The sidecar shape (`<slide>.analysis.json`) is
documented in the brief that introduced this module - see workroom's `/api/analysis` handler.
"""
import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

WHOLE_SHEET_MAX_ROWS = 400
WHOLE_SHEET_MAX_COLS = 30

_A1_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")

# A cell/range qualified by a sheet name: 'Sheet Name'!D6, Sheet1!D6:D8, etc.
_QUALIFIED_RE = re.compile(
    r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_. ]*))!"
    r"\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?"
)
# A same-sheet cell/range: D6, $D$6, D6:D8. Bounded so it never matches inside a function name
# (letters immediately followed by digits is what a ref looks like; a function name like SUM(
# has no digits touching the letters) and never re-matches what _QUALIFIED_RE already consumed
# (that text is blanked out before this pattern runs).
_PLAIN_RE = re.compile(
    r"(?<![A-Za-z0-9_])\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?(?![A-Za-z0-9(_])"
)

_WB_CACHE = {}   # path str -> (mtime, formulas_workbook, values_workbook)


# ---------------------------------------------------------------- column <-> index helpers

def _col_to_idx(col):
    idx = 0
    for ch in col.upper():
        idx = idx * 26 + (ord(ch) - 64)
    return idx


def _idx_to_col(idx):
    s = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def _parse_a1(ref):
    m = _A1_RE.match(ref.strip())
    if not m:
        raise ValueError("bad cell reference: %r" % ref)
    return _col_to_idx(m.group(1)), int(m.group(2))


def _parse_range(range_a1):
    parts = range_a1.split(":")
    c1, r1 = _parse_a1(parts[0])
    if len(parts) == 2:
        c2, r2 = _parse_a1(parts[1])
    else:
        c2, r2 = c1, r1
    return min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)


def _expand_range(c1, r1, c2, r2):
    i1, i2 = _col_to_idx(c1), _col_to_idx(c2)
    out = []
    for r in range(min(r1, r2), max(r1, r2) + 1):
        for ci in range(min(i1, i2), max(i1, i2) + 1):
            out.append(_idx_to_col(ci) + str(r))
    return out


# ---------------------------------------------------------------- workbook cache

def _get_workbooks(path):
    """Return (formulas_workbook, values_workbook) for `path`, cached by (path, mtime)."""
    path = Path(path)
    mtime = path.stat().st_mtime
    key = str(path)
    cached = _WB_CACHE.get(key)
    if cached and cached[0] == mtime:
        return cached[1], cached[2]
    if cached:
        for wb in (cached[1], cached[2]):
            try:
                wb.close()
            except Exception:  # noqa: BLE001 - never let a stale-handle close fail a read
                pass
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=True)
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=True)
    _WB_CACHE[key] = (mtime, wb_f, wb_v)
    return wb_f, wb_v


# ---------------------------------------------------------------- cell/grid building

def _cell_value(v_cell):
    v = v_cell.value
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def _cell_info(f_cell, v_cell, ref):
    """Build one cell's info dict. `ref` is passed in (computed from the loop position) rather
    than read off the cell, because a read_only worksheet hands back a bare EmptyCell - no
    .row/.column/.coordinate/.font - for any cell that has no data at all, which a used range
    like a databook sheet's blank separator rows hits constantly."""
    raw_f = f_cell.value
    is_formula = isinstance(raw_f, str) and raw_f.startswith("=")
    formula = raw_f if is_formula else None
    value = _cell_value(v_cell)

    if is_formula:
        kind = "link" if "!" in raw_f else "formula"
    elif value is None:
        kind = "empty"
    elif isinstance(value, (int, float)):
        kind = "input"
    else:
        kind = "text"

    if isinstance(value, (int, float)):
        out_value = value
    elif value is None:
        out_value = None
    else:
        out_value = str(value)

    try:
        bold = bool(f_cell.font and f_cell.font.bold)
    except Exception:  # noqa: BLE001 - formatting is a nicety, never fatal to a read
        bold = False

    return {"ref": ref, "value": out_value, "formula": formula, "kind": kind, "bold": bold}


def _build_grid(ws_f, ws_v, min_row, max_row, min_col, max_col):
    cols = [_idx_to_col(c) for c in range(min_col, max_col + 1)]
    rows = []
    f_iter = ws_f.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
    v_iter = ws_v.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
    for row_num, (f_row, v_row) in enumerate(zip(f_iter, v_iter), start=min_row):
        if not f_row:
            continue
        cells = [_cell_info(fc, vc, col + str(row_num))
                 for col, fc, vc in zip(cols, f_row, v_row)]
        rows.append({"row": row_num, "cells": cells})
    return {
        "rows": rows,
        "cols": cols,
        "min_row": min_row,
        "max_row": max_row,
        "min_col": cols[0] if cols else None,
        "max_col": cols[-1] if cols else None,
    }


# ---------------------------------------------------------------- public API

def read_range(xlsx_path, sheet, range_a1, context_rows=2):
    """Grid for `range_a1` on `sheet` in `xlsx_path`, expanded by `context_rows` above and
    below and clamped to the sheet's used area, plus `sheet_used` (max row/col) so the caller
    can offer a "whole sheet" view. Never raises: a missing file or sheet comes back as
    {"error": "..."}."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        return {"error": "file not found: %s" % xlsx_path}
    try:
        wb_f, wb_v = _get_workbooks(xlsx_path)
    except Exception as e:  # noqa: BLE001 - a corrupt/locked file is a data issue, not a crash
        return {"error": "could not open workbook: %s: %s" % (type(e).__name__, e)}
    if sheet not in wb_f.sheetnames:
        return {"error": "sheet not found: %s" % sheet}

    try:
        min_col, min_row, max_col, max_row = _parse_range(range_a1)
    except ValueError as e:
        return {"error": str(e)}

    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet]
    sheet_max_row = ws_f.max_row or 1
    sheet_max_col = ws_f.max_column or 1

    min_row = max(1, min_row - context_rows)
    max_row = min(sheet_max_row, max_row + context_rows)
    min_col = max(1, min_col)
    max_col = min(sheet_max_col, max_col)

    grid = _build_grid(ws_f, ws_v, min_row, max_row, min_col, max_col)
    grid["sheet_used"] = {"max_row": sheet_max_row, "max_col": sheet_max_col}
    return grid


def whole_sheet(xlsx_path, sheet):
    """The full sheet (capped at WHOLE_SHEET_MAX_ROWS x WHOLE_SHEET_MAX_COLS) for the "whole
    sheet" toggle. Same error shape as read_range."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        return {"error": "file not found: %s" % xlsx_path}
    try:
        wb_f, wb_v = _get_workbooks(xlsx_path)
    except Exception as e:  # noqa: BLE001
        return {"error": "could not open workbook: %s: %s" % (type(e).__name__, e)}
    if sheet not in wb_f.sheetnames:
        return {"error": "sheet not found: %s" % sheet}

    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet]
    sheet_max_row = ws_f.max_row or 1
    sheet_max_col = ws_f.max_column or 1
    max_row = min(sheet_max_row, WHOLE_SHEET_MAX_ROWS)
    max_col = min(sheet_max_col, WHOLE_SHEET_MAX_COLS)

    grid = _build_grid(ws_f, ws_v, 1, max_row, 1, max_col)
    grid["sheet_used"] = {"max_row": sheet_max_row, "max_col": sheet_max_col}
    return grid


def precedents(formula, sheet=None):
    """Same-sheet cell references a formula reads from (ranges expanded, `$` stripped), plus
    any cross-sheet references as `links` [{"sheet": ..., "ref": ...}] - a ref qualified by
    `sheet` itself (the formula's own sheet) counts as same-sheet, not a link. Returns
    {"refs": [...], "links": [...]}; never raises on unparseable input."""
    if not formula:
        return {"refs": [], "links": []}

    same_sheet_refs = []
    links = []

    def qual_sub(m):
        sh = m.group(1) or m.group(2)
        c1, r1 = m.group(3), int(m.group(4))
        if m.group(5):
            c2, r2 = m.group(5), int(m.group(6))
        else:
            c2, r2 = c1, r1
        expanded = _expand_range(c1, r1, c2, r2)
        if sheet is not None and sh == sheet:
            same_sheet_refs.extend(expanded)
        else:
            for r in expanded:
                links.append({"sheet": sh, "ref": r})
        return " "

    remainder = _QUALIFIED_RE.sub(qual_sub, formula)

    refs = list(same_sheet_refs)

    def plain_sub(m):
        c1, r1 = m.group(1), int(m.group(2))
        if m.group(3):
            c2, r2 = m.group(3), int(m.group(4))
        else:
            c2, r2 = c1, r1
        refs.extend(_expand_range(c1, r1, c2, r2))
        return " "

    _PLAIN_RE.sub(plain_sub, remainder)

    seen = set()
    out_refs = []
    for r in refs:
        if r not in seen:
            seen.add(r)
            out_refs.append(r)

    seen_links = set()
    out_links = []
    for link in links:
        key = (link["sheet"], link["ref"])
        if key not in seen_links:
            seen_links.add(key)
            out_links.append(link)

    return {"refs": out_refs, "links": out_links}


def load_sidecar(deck_dir, slide_file):
    """`<deck_dir>/<slide stem>.analysis.json` parsed, or None if it doesn't exist or isn't
    valid JSON. Never raises."""
    deck_dir = Path(deck_dir)
    p = deck_dir / (Path(slide_file).stem + ".analysis.json")
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
