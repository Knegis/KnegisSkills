"""Translate extracted HTML-slide geometry (rects + text boxes) into a native
PPT slide on the configured template.

Usage: python html_to_pptx.py <geom.json> <out.pptx>

Canvas: 1280x720 px -> 12192000x6858000 EMU.  px*9525 = EMU ;  px*0.75 = pt.
Avoids the known corruptors: unique cNvPr ids, srgbClr (never an invalid enum),
real <a:p>/<a:r> structure, anchor on bodyPr, pack via office_io (validated by python-pptx).
"""
import sys, json, re, shutil, html, hashlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
EMU = 9525                      # per css px
import _cli                                 # noqa: E402 - shared --help / arg-count handling
import config as _cfg                       # house settings + machine paths; see scripts/config.py
import office_io
import patch_fonts as _pf
TPL = _cfg.require("the template .pptx", _cfg.TEMPLATE)
FONT = _cfg.FONT_NAME
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Layouts are resolved by NAME against whatever template is configured, so the skill works on
# any 16:9 deck. An HTML slide opts into a layout via <body data-ppt-layout="...">: an alias
# below, a substring of a layout name (case-insensitive), or a raw slideLayoutN.xml.
# Default = the blankest layout available ("Blank", "No Content", ... - fewest placeholders).
LAYOUT_ALIASES = {
    "blank":   ["blank", "no content", "empty"],
    "title":   ["title slide", "title -", "title page", "cover"],
    "divider": ["section header", "divider", "section"],
    "end":     ["end slide", "closing", "thank"],
}

def _layouts(unp):
    """[(filename, name, n_placeholders)] for the unpacked template."""
    out = []
    for lf in sorted((Path(unp) / "ppt" / "slideLayouts").glob("slideLayout*.xml"),
                     key=lambda p: int(re.search(r"(\d+)", p.name).group(1))):
        x = lf.read_text(encoding="utf-8", errors="ignore")
        m = re.search(r'<p:cSld[^>]*\bname="([^"]*)"', x)
        out.append((lf.name, (m.group(1) if m else ""), len(re.findall(r"<p:ph\b", x))))
    return out

def _pick(hits, lays, v):
    """When several layouts match, prefer one whose name contains 'white'/'light'; otherwise
    take the first and say so, so an ambiguous match is never silent."""
    if len(hits) == 1:
        return hits[0]
    by_name = {f: name for f, name, _ in lays}
    preferred = [f for f in hits if re.search(r"white|light", by_name.get(f, ""), re.I)]
    if len(preferred) == 1:
        return preferred[0]          # the white/light preference settles it - nothing to report
    chosen = preferred[0] if preferred else hits[0]
    print(f"  ! layout '{v}' matched {len(hits)} layouts, using '{by_name.get(chosen, chosen)}' "
          f"(candidates: {', '.join(by_name.get(f, f) for f in hits)})")
    return chosen

def layout_target(data, unp):
    # Resolution order: explicit data-ppt-layout (filename / alias / substring) -> the house
    # default_layout as a substring match -> alias "blank" -> the blankest layout (fewest
    # placeholders). See config.DEFAULT_LAYOUT (house.json "default_layout").
    v = (data.get("layout") or "").strip()
    lays = _layouts(unp)
    if not lays:
        raise SystemExit("template has no slide layouts")
    if re.match(r"^slideLayout\d+\.xml$", v) and any(f == v for f, _, _ in lays):
        return v
    if v:
        wanted = LAYOUT_ALIASES.get(v.lower(), [v.lower()])
        for w in wanted:
            hits = [f for f, name, _ in lays if w in name.lower()]
            if hits:
                return _pick(hits, lays, v)
        print(f"  ! layout '{v}' not found in template - trying the house default layout")
    default = (_cfg.DEFAULT_LAYOUT or "").strip()
    if default:
        hits = [f for f, name, _ in lays if default.lower() in name.lower()]
        if hits:
            return _pick(hits, lays, default)
        print(f"  ! default_layout '{default}' not found in template - using the blankest layout")
    for w in LAYOUT_ALIASES["blank"]:
        hits = [f for f, name, _ in lays if w in name.lower()]
        if hits:
            return _pick(hits, lays, w)
    return min(lays, key=lambda t: t[2])[0]

def px(v): return int(round(v * EMU))
SHRINK = 0.94                                      # absorb PPT vs browser metric drift
def half_pt(p): return round(p * 2) / 2.0          # snap a pt value to nearest 0.5pt
def pt100(v):                                       # css px -> pt*100 (centipoints)
    # House rule: font/spacing sizes are only ever whole points or .5 - never 13.7, 7.8 etc.
    return int(round(v * 0.75 * SHRINK * 2)) * 50   # snap to nearest 0.5pt, in centipoints

def color(css):
    """rgb()/rgba()/#hex/transparent -> (HEX, alpha 0..100000 or None)."""
    if not css: return None, None
    css = css.strip()
    if css.lower() == "transparent": return None, None
    h = re.fullmatch(r'#([0-9a-fA-F]{3}|[0-9a-fA-F]{6})', css)
    if h:
        v = h.group(1)
        if len(v) == 3: v = "".join(c * 2 for c in v)
        return v.upper(), None
    m = re.match(r'rgba?\(([^)]+)\)', css)
    if not m: return None, None
    parts = [p.strip() for p in m.group(1).split(',')]
    r, g, b = (int(float(parts[i])) for i in range(3))
    a = float(parts[3]) if len(parts) > 3 else 1.0
    return "%02X%02X%02X" % (r, g, b), (None if a >= 0.999 else int(a*100000))

def grad_top_color(grad):
    cols = re.findall(r'rgba?\([^)]+\)', grad)
    return color(cols[-1])[0] if cols else None

class IdGen:
    def __init__(self): self.n = 1
    def __call__(self): self.n += 1; return self.n

CANVAS_W, CANVAS_H = 1280, 720   # the HTML canvas that maps 1:1 to a 16:9 slide

IMG_CONTENT_TYPES = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                      "gif": "image/gif", "bmp": "image/bmp", "webp": "image/webp"}

class MediaRegistry:
    """De-dupes embedded image files by content hash across a whole unpacked .pptx
    package (writes them under ppt/media/), and tracks which [Content_Types].xml
    Default extensions the package needs. One instance per build - shared across every
    slide of a deck (build_deck.py) so the same photo used on two slides is only stored
    once - and call ensure_content_types() once, after all slides are built."""
    def __init__(self, unp):
        self.unp = Path(unp)
        self.media_dir = self.unp / "ppt" / "media"
        self.media_dir.mkdir(parents=True, exist_ok=True)
        self.by_hash = {}      # sha256 -> media filename already written (e.g. "image3.png")
        self.exts_used = set()
        self.n = 0

    def add(self, src_path):
        """Copy src_path into ppt/media/ (unless an identical file is already there) and
        return the media filename to reference from a slide's blip relationship."""
        data = Path(src_path).read_bytes()
        h = hashlib.sha256(data).hexdigest()
        cached = self.by_hash.get(h)
        if cached:
            return cached
        ext = Path(src_path).suffix.lower().lstrip(".")
        if ext not in IMG_CONTENT_TYPES:
            ext = "png"                     # anything already rasterised (e.g. from SVG) lands here
        # The template being built on may already ship media parts (logos, theme images) -
        # never reuse one of those names, or a same-numbered part with a different
        # extension would silently overwrite unrelated template media.
        while True:
            self.n += 1
            fname = "image%d.%s" % (self.n, ext)
            if not (self.media_dir / fname).exists():
                break
        (self.media_dir / fname).write_bytes(data)
        self.by_hash[h] = fname
        self.exts_used.add(ext)
        return fname

    def ensure_content_types(self):
        ct = self.unp / "[Content_Types].xml"
        s = ct.read_text(encoding="utf-8")
        for ext in sorted(self.exts_used):
            if ('Extension="%s"' % ext) in s:
                continue
            ctype = IMG_CONTENT_TYPES.get(ext, "image/" + ext)
            s = s.replace("</Types>", '<Default Extension="%s" ContentType="%s"/></Types>' % (ext, ctype))
        ct.write_text(s, encoding="utf-8")

class SlideRels:
    """Per-slide relationship-id allocator for images. rId1 is reserved for the
    slideLayout relationship (written separately by the caller); image rIds start at 2."""
    def __init__(self, start=2):
        self.next_id = start
        self.entries = []       # [(rId, media filename)]

    def add_image(self, media_filename):
        rid = "rId%d" % self.next_id
        self.next_id += 1
        self.entries.append((rid, media_filename))
        return rid

    def rels_xml(self):
        return "".join('<Relationship Id="%s" Type="%s/image" Target="../media/%s"/>' % (rid, RT, fname)
                        for rid, fname in self.entries)

def _cover_crop(natural_w, natural_h, box_w, box_h, pos_x, pos_y, scale=None):
    """Fractions (0..1) to crop from (left, top, right, bottom) of a natural_w x natural_h
    image so that, once cropped, it exactly fills box_w x box_h with no distortion -
    the "cover" fit. `scale=None` computes the cover scale (max of the two axis ratios);
    passing scale=1.0 gives the "none" fit's crop (natural size, cropped only where it
    overflows the box). pos_x/pos_y (0..1) place the crop window within the excess."""
    if natural_w <= 0 or natural_h <= 0 or box_w <= 0 or box_h <= 0:
        return 0.0, 0.0, 0.0, 0.0
    if scale is None:
        scale = max(box_w / natural_w, box_h / natural_h)
    scaled_w, scaled_h = natural_w * scale, natural_h * scale
    excess_w = max(0.0, scaled_w - box_w)
    excess_h = max(0.0, scaled_h - box_h)
    l = (pos_x * excess_w / scaled_w) if scaled_w else 0.0
    r = ((1 - pos_x) * excess_w / scaled_w) if scaled_w else 0.0
    t = (pos_y * excess_h / scaled_h) if scaled_h else 0.0
    b = ((1 - pos_y) * excess_h / scaled_h) if scaled_h else 0.0
    return l, t, r, b

def _contain_box(natural_w, natural_h, box_x, box_y, box_w, box_h, pos_x, pos_y, only_shrink=False):
    """New (x, y, w, h) for a picture letterboxed to fit inside box_x/y/w/h without
    cropping - the "contain" (and "scale-down", only_shrink=True) fit."""
    if natural_w <= 0 or natural_h <= 0 or box_w <= 0 or box_h <= 0:
        return box_x, box_y, box_w, box_h
    scale = min(box_w / natural_w, box_h / natural_h)
    if only_shrink:
        scale = min(scale, 1.0)
    new_w, new_h = natural_w * scale, natural_h * scale
    new_x = box_x + (box_w - new_w) * pos_x
    new_y = box_y + (box_h - new_h) * pos_y
    return new_x, new_y, new_w, new_h

def _src_rect_xml(l, t, r, b):
    def pct(frac): return int(round(frac * 100000))
    return '<a:srcRect l="%d" t="%d" r="%d" b="%d"/>' % (pct(l), pct(t), pct(r), pct(b))

def pic_shape(idg, im, media, slide_rels):
    fname = media.add(im["src"])
    rid = slide_rels.add_image(fname)

    box_x, box_y, box_w, box_h = im["x"], im["y"], im["w"], im["h"]
    nat_w, nat_h = im.get("natural_w") or 0, im.get("natural_h") or 0
    pos_x, pos_y = im.get("pos_x", 0.5), im.get("pos_y", 0.5)
    fit = im.get("fit") or "fill"

    src_rect_xml = ""
    if fit == "cover" and nat_w and nat_h:
        src_rect_xml = _src_rect_xml(*_cover_crop(nat_w, nat_h, box_w, box_h, pos_x, pos_y))
    elif fit in ("contain", "scale-down") and nat_w and nat_h:
        box_x, box_y, box_w, box_h = _contain_box(nat_w, nat_h, box_x, box_y, box_w, box_h,
                                                    pos_x, pos_y, only_shrink=(fit == "scale-down"))
    elif fit == "none" and nat_w and nat_h:
        src_rect_xml = _src_rect_xml(*_cover_crop(nat_w, nat_h, box_w, box_h, pos_x, pos_y, scale=1.0))
    # fit == "fill" (or no natural size to work from): stretch to the box, no crop.

    x, y, w, h = px(box_x), px(box_y), px(box_w), px(box_h)
    radius = im.get("radius") or 0
    short_px = min(im["w"], im["h"])        # circle test against the ELEMENT box, not a shrunk pic box
    geom = "ellipse" if (radius and radius >= short_px * 0.49) else "rect"

    opacity = im.get("opacity")
    if opacity is None:
        opacity = 1.0
    blip_xml = ('<a:blip r:embed="%s"><a:alphaModFix amt="%d"/></a:blip>' % (rid, int(round(opacity * 100000)))
                if opacity < 0.999 else '<a:blip r:embed="%s"/>' % rid)

    return ('<p:pic><p:nvPicPr><p:cNvPr id="%d" name="Picture %d"/>'
            '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
            '<p:blipFill>%s%s<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="%s"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
            % (idg(), idg.n, blip_xml, src_rect_xml, x, y, w, h, geom))

def fill_xml(hexv, alpha):
    if hexv is None: return "<a:noFill/>"
    inner = '<a:srgbClr val="%s"%s/>' % (hexv, "" if alpha is None else ("><a:alpha val=\"%d\"/></a:srgbClr>" % alpha))
    if alpha is None:
        return "<a:solidFill><a:srgbClr val=\"%s\"/></a:solidFill>" % hexv
    return "<a:solidFill><a:srgbClr val=\"%s\"><a:alpha val=\"%d\"/></a:srgbClr></a:solidFill>" % (hexv, alpha)

def _is_thin_rule(r):
    """A filled div this thin, with no border/gradient/radius, reads as a rule, not a
    rectangle shape - a user pinned exactly this on a real deck ("this is turned into a
    rectangle shape in the ppt. I'd prefer it to be a line object"). Qualifies as either a
    horizontal rule (h<=8px) or a vertical one (w<=8px); a near-square dot (w==h) stays a
    rect, since it isn't "horizontal or vertical" at all."""
    if not r.get("fill") or r.get("border") or r.get("grad") or r.get("radius"):
        return False
    w, h = r["w"], r["h"]
    if w == h:
        return False
    return h <= 8 or w <= 8

def line_shape(idg, r):
    """A thin filled rect emitted as a straight connector (<p:cxnSp>) along the centre
    line of the original box, with the box's short side as the line's stroke weight -
    so PowerPoint shows a line object, not a flat rectangle. Orientation follows the
    long axis: wider-than-tall -> horizontal line, taller-than-wide -> vertical line."""
    hexv, alpha = color(r["fill"])
    x, y, w, h = r["x"], r["y"], r["w"], r["h"]
    horizontal = w > h
    if horizontal:
        thickness = h
        off_x, off_y = px(x), px(y + h / 2.0)
        ext_cx, ext_cy = px(w), 0
    else:
        thickness = w
        off_x, off_y = px(x + w / 2.0), px(y)
        ext_cx, ext_cy = 0, px(h)
    return ('<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="%d" name="Straight Connector %d"/>'
            '<p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
            '<a:ln w="%d" cap="flat">%s</a:ln></p:spPr></p:cxnSp>'
            % (idg(), idg.n, off_x, off_y, ext_cx, ext_cy, px(thickness), fill_xml(hexv, alpha)))

def rect_shape(idg, r):
    if _is_thin_rule(r):
        return line_shape(idg, r)
    x, y, w, h = px(r["x"]), px(r["y"]), px(r["w"]), px(r["h"])
    hexv, alpha = color(r["fill"]) if r["fill"] else (None, None)
    if hexv is None and r["grad"]:
        hexv, alpha = grad_top_color(r["grad"]), None
    radius = r.get("radius", 0)
    # A border-radius >= ~half the shorter side (e.g. CSS border-radius:50% on an
    # equal-sided box) is a circle/ellipse, NOT a rounded rect - a capped roundRect
    # renders as an ugly stadium/squircle. Emit a real ellipse so circles stay circular.
    short_px = min(r["w"], r["h"])
    if radius and radius >= short_px * 0.49:
        geom = "ellipse"
    elif radius and radius > 1:
        geom = "roundRect"
    else:
        geom = "rect"
    shape_name = {"ellipse": "Oval", "roundRect": "Rounded Rectangle"}.get(geom, "Rectangle")
    adj = ""
    if geom == "roundRect":
        # adj fraction of shorter side, capped
        frac = min(0.5, (radius * EMU) / max(1, min(w, h)))
        adj = '<a:avLst><a:gd name="adj" fmla="val %d"/></a:avLst>' % int(frac*100000)
    else:
        adj = "<a:avLst/>"
    # border: emit a single outline if any side present (uses top side spec)
    ln = ""
    b = r["border"]
    if b:
        side_w = max(b["t"], b["l"], b["r"], b["b"])
        # if only one side (e.g. left accent or open bracket), still approximate with full outline
        bc, ba = color(b["colT"] or b["colL"] or b["colR"] or b["colB"])
        if bc and side_w > 0:
            # honour rgba alpha on borders too - a faint rgba ring must stay faint, not
            # render as a harsh solid line (seen on a real deck: rgba ring borders came out solid white).
            inner = ('<a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr>' % (bc, ba)) if ba is not None else ('<a:srgbClr val="%s"/>' % bc)
            ln = '<a:ln w="%d"><a:solidFill>%s</a:solidFill></a:ln>' % (px(side_w), inner)
    return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s %d"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="%s">%s</a:prstGeom>%s%s</p:spPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
            % (idg(), shape_name, idg.n, x, y, w, h, geom, adj, fill_xml(hexv, alpha), ln))

def run_xml(run):
    t = run["t"].upper() if run.get("upper") else run["t"]
    t = html.escape(t)
    hexv, _ = color(run["color"])
    b = ' b="1"' if run["weight"] >= 600 else ""
    i = ' i="1"' if run["italic"] else ""
    spc = (' spc="%d"' % int(run["ls"]*0.75*100)) if run.get("ls") else ""
    sz = pt100(run["size"])
    return ('<a:r><a:rPr lang="en-US" sz="%d"%s%s%s><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:latin typeface="%s"/></a:rPr><a:t>%s</a:t></a:r>'
            % (sz, b, i, spc, hexv or "000000", FONT, t))

ALIGN = {"left":"l","center":"ctr","right":"r","start":"l","end":"r","justify":"just"}

def _txsp(idg, x, y, w, h, bodypr, paras):
    return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="TextBox %d"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody>%s<a:lstStyle/>%s</p:txBody></p:sp>'
            % (idg(), idg.n, x, y, w, h, bodypr, paras))

def text_shape(idg, t):
    x, y, w, h = px(t["x"]), px(t["y"]), px(t["w"]), px(t["h"])
    anchor = t.get("anchor", "t")
    # grouped: one text box, child lines -> paragraphs with measured spacing
    # (space-before in pt + line spacing %), so PPT recreates the exact gaps.
    if t.get("grouped"):
        paras = ""
        for gp in t["gparas"]:
            algn = ALIGN.get(gp.get("align") or t["align"], "l")
            spc = ""
            # Exact line spacing in pt, not a percentage: PPT's "100%" is its own single
            # spacing (~1.15-1.2 em for most fonts), so a CSS line-height of 1.5 emitted as
            # 150% renders ~1.7 em and a list that fit in the browser overflows its box.
            if gp.get("lhpx"):
                spc += '<a:lnSpc><a:spcPts val="%d"/></a:lnSpc>' % pt100(gp["lhpx"])
            elif gp.get("lhpct"):
                spc += '<a:lnSpc><a:spcPct val="%d"/></a:lnSpc>' % gp["lhpct"]
            if (gp.get("spcBef") or 0) > 0.3:
                spc += '<a:spcBef><a:spcPts val="%d"/></a:spcBef>' % pt100(gp["spcBef"])
            runs = "".join(run_xml(r) for r in gp["runs"])
            paras += '<a:p><a:pPr algn="%s">%s</a:pPr>%s</a:p>' % (algn, spc, runs)
        bodypr = ('<a:bodyPr wrap="square" anchor="%s" lIns="%d" tIns="%d" rIns="%d" bIns="%d"/>'
                  % (anchor, px(t["pl"]), px(t["pt"]), px(t["pr"]), px(t["pb"])))
        return _txsp(idg, x, y, w, h, bodypr, paras)
    algn = ALIGN.get(t["align"], "l")
    lines = t.get("lines", 1)
    # Multi-line text keeps the browser's line pitch exactly (see the note on grouped text).
    lnspc = ('<a:lnSpc><a:spcPts val="%d"/></a:lnSpc>' % pt100(t["lhpx"])) if (lines > 1 and t.get("lhpx")) else ""
    paras = ""
    for p in t["paras"]:
        runs = "".join(run_xml(r) for r in p)
        paras += ('<a:p><a:pPr algn="%s">%s</a:pPr>%s</a:p>' % (algn, lnspc, runs)) if lnspc                  else ('<a:p><a:pPr algn="%s"/>%s</a:p>' % (algn, runs))
    # single short label that the browser rendered on ONE line -> don't wrap (prevents ugly
    # mid-word breaks of tags/wordmark/chips from metric drift). If the browser already broke
    # it onto two lines, it must stay wrapped in PPT too, or it rides out of its box.
    nruns = sum(len(p) for p in t["paras"])
    total_len = sum(len(r["t"]) for p in t["paras"] for r in p)
    nowrap = (len(t["paras"]) == 1 and total_len <= 34 and lines <= 1)
    wrap = "none" if nowrap else "square"
    bodypr = ('<a:bodyPr wrap="%s" anchor="%s" lIns="%d" tIns="%d" rIns="%d" bIns="%d"/>'
              % (wrap, anchor, px(t["pl"]), px(t["pt"]), px(t["pr"]), px(t["pb"])))
    return _txsp(idg, x, y, w, h, bodypr, paras)

def build_slide_xml(data, idg, media=None, slide_rels=None):
    shapes = []
    # Background comes from the chosen layout (house.json default_layout, e.g. "Blank").
    # A non-white HTML body colour becomes the slide's REAL background (<p:bg>), never a
    # full-bleed rectangle shape: a rect is a selectable object that sits in the z-order,
    # traps clicks, shows up in the selection pane and gets dragged by accident ("a random
    # background page", as a reviewer put it). Slide background = slide property.
    bg_hex, bg_alpha = (color(data["bodyBg"]) if data.get("bodyBg") else (None, None))
    if not (bg_hex and bg_hex.upper() != "FFFFFF"):
        bg_hex, bg_alpha = None, None

    # Same promotion for a full-bleed DIV at the bottom of the stack - the usual case, because
    # slide canvases are authored as `.slide{width:1280px;height:720px;background:<colour>}`.
    # Without this it lands in the .pptx as a giant beige rectangle covering the slide.
    # Only opaque, square, borderless, canvas-sized rects at the BOTTOM of the z-order are
    # promoted; a later full-bleed rect is a deliberate scrim/overlay and stays a shape.
    rects = list(data["rects"])
    while rects:
        r = rects[0]
        if not (r["fill"] and not r["border"] and not r.get("radius")
                and r["x"] <= 0.5 and r["y"] <= 0.5
                and r["w"] >= CANVAS_W - 1 and r["h"] >= CANVAS_H - 1):
            break
        hexv, alph = color(r["fill"])
        if hexv is None or alph is not None:      # transparent / scrim - leave it as a shape
            break
        bg_hex, bg_alpha = hexv, None
        rects.pop(0)

    bg_xml = ""
    if bg_hex:
        bg_xml = "<p:bg><p:bgPr>%s<a:effectLst/></p:bgPr></p:bg>" % fill_xml(bg_hex, bg_alpha)
    # Pictures interleave with rects and texts by document (z) order - a full-bleed photo
    # under a scrim rect under white title text is the primary use case (cover slides).
    # Hand-written geometry (e.g. in unit tests) may omit "z"; default rects to the very
    # bottom and texts to the very top so the old "all rects, then all texts" order still
    # holds for that data, unchanged.
    items = [(r.get("z", -1), "rect", r) for r in rects]
    items += [(t.get("z", 10**9), "text", t) for t in data["texts"]]
    items += [(im.get("z", 0), "image", im) for im in (data.get("images") or [])]
    items.sort(key=lambda it: it[0])
    for _, kind, obj in items:
        if kind == "rect":
            shapes.append(rect_shape(idg, obj))
        elif kind == "text":
            shapes.append(text_shape(idg, obj))
        elif media is not None and slide_rels is not None:
            shapes.append(pic_shape(idg, obj, media, slide_rels))
    body = "".join(shapes)
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:cSld>%s<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
            '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
            '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
            '%s</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>' % (bg_xml, body))

def register(unp):
    ct = unp/"[Content_Types].xml"; s = ct.read_text(encoding="utf-8")
    if "/ppt/slides/slide1.xml" not in s:
        s = s.replace("</Types>", '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/></Types>')
        ct.write_text(s, encoding="utf-8")
    pr = unp/"ppt"/"_rels"/"presentation.xml.rels"; s = pr.read_text(encoding="utf-8")
    rids = [int(m) for m in re.findall(r'Id="rId(\d+)"', s)]; rid = "rId%d" % (max(rids)+1 if rids else 1)
    s = s.replace("</Relationships>", '<Relationship Id="%s" Type="%s/slide" Target="slides/slide1.xml"/></Relationships>' % (rid, RT))
    pr.write_text(s, encoding="utf-8")
    pp = unp/"ppt"/"presentation.xml"; s = pp.read_text(encoding="utf-8")
    sid = max([int(i) for i in re.findall(r'<p:sldId[^>]*id="(\d+)"', s)] or [255]) + 1
    entry = '<p:sldId id="%d" r:id="%s"/>' % (sid, rid)
    if "<p:sldIdLst" in s:
        s = s.replace("</p:sldIdLst>", entry + "</p:sldIdLst>")
    elif "</p:handoutMasterIdLst>" in s:
        s = s.replace("</p:handoutMasterIdLst>", "</p:handoutMasterIdLst><p:sldIdLst>%s</p:sldIdLst>" % entry)
    else:
        s = s.replace("</p:sldMasterIdLst>", "</p:sldMasterIdLst><p:sldIdLst>%s</p:sldIdLst>" % entry)
    pp.write_text(s, encoding="utf-8")

CHART_TYPES = {"stacked-column":"COLUMN_STACKED","column":"COLUMN_CLUSTERED",
               "stacked-bar":"BAR_STACKED","bar":"BAR_CLUSTERED","line":"LINE","line-markers":"LINE_MARKERS"}
LABEL_POS = {"center":"CENTER","inside_end":"INSIDE_END","inside_base":"INSIDE_BASE",
             "outside_end":"OUTSIDE_END","above":"ABOVE"}
# PowerPoint only accepts ctr / inBase / inEnd on STACKED charts. python-pptx will happily
# write outEnd or above, and python-pptx will read the file back fine - but PowerPoint
# refuses to open it at all ("corrupted and unreadable"). Silently clamp to a legal value.
STACKED = {"COLUMN_STACKED", "BAR_STACKED"}
STACKED_LEGAL = {"CENTER", "INSIDE_END", "INSIDE_BASE"}

def add_charts_to_slide(slide, charts):
    """Insert real, editable PPT charts (python-pptx) at each .ppt-chart box on ONE slide.
    Hybrid: native chart owns the data viz; the rest of the slide stays as translated shapes."""
    from pptx.util import Emu, Pt
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.dml.color import RGBColor
    for ch in charts:
        cd = CategoryChartData()
        cd.categories = ch["categories"]
        for s in ch["series"]:
            cd.add_series(s["name"], tuple(s["values"]))
        x, y, cw, cyh = (Emu(int(round(ch[k] * EMU))) for k in ("x", "y", "w", "h"))
        ctype_name = CHART_TYPES.get(ch.get("type"), "COLUMN_STACKED")
        ctype = getattr(XL_CHART_TYPE, ctype_name)
        chart = slide.shapes.add_chart(ctype, x, y, cw, cyh, cd).chart
        chart.font.name = FONT; chart.font.size = Pt(11); chart.has_title = False
        if ctype_name in ("COLUMN_STACKED", "COLUMN_CLUSTERED", "BAR_STACKED", "BAR_CLUSTERED"):
            # chart_preview.js draws each bar at 56% of its category slot (22% pad each side);
            # PowerPoint's default gap width of 150 draws them far thinner. 79 = 44/56 matches.
            try: chart.plots[0].gap_width = 79
            except Exception: pass
        if ctype_name in ("BAR_CLUSTERED", "BAR_STACKED"):
            # Horizontal bar chart: the HTML preview (chart_preview.js) always draws the
            # first category at the top, but PowerPoint's own default puts it at the
            # bottom - reverse the category axis so PowerPoint matches the preview.
            # NOTE: also pinning value_axis.crosses to MAXIMUM (as ECMA-376 might
            # suggest for keeping the crossing point put) was tried and rendered
            # verifiably wrong in PowerPoint - it flips the value scale (bars grow
            # from the right, in reverse magnitude order) and drops the bar fill.
            # reverse_order alone renders correctly: bars keep growing left-to-right.
            try:
                chart.category_axis.reverse_order = True
            except Exception: pass
        if ch.get("legend", "top") == "none":
            chart.has_legend = False
        else:
            chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False; chart.legend.font.size = Pt(10)
        for si, s in enumerate(ch["series"]):
            ser = chart.series[si]
            # PowerPoint's default invertIfNegative draws negative bars as white hollow boxes
            try: ser.invert_if_negative = False
            except Exception: pass
            if s.get("color"):
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(s["color"])
            labels = s.get("labels")
            if labels:
                ser.has_data_labels = True
                dl = ser.data_labels
                # has_data_labels alone shows nothing - PowerPoint needs an explicit
                # show_value (or per-point custom text below) or the labels stay blank.
                if not isinstance(labels, list):
                    dl.show_value = True
                dl.font.size = Pt(half_pt(s.get("labelsize", 9))); dl.font.name = FONT
                if s.get("labelbold"): dl.font.bold = True
                if s.get("labelcolor"): dl.font.color.rgb = RGBColor.from_string(s["labelcolor"])
                if s.get("labelpos"):
                    want = LABEL_POS.get(s["labelpos"], "CENTER")
                    if ctype_name in STACKED and want not in STACKED_LEGAL:
                        print(f"  ! labelpos '{s['labelpos']}' is illegal on a stacked chart "
                              f"(PowerPoint would reject the file) - using inside_end instead")
                        want = "INSIDE_END"
                    try: dl.position = getattr(XL_LABEL_POSITION, want)
                    except Exception: pass
                else:
                    want = None
                if isinstance(labels, list):
                    for pi, txt in enumerate(labels):
                        try:
                            pdl = ser.points[pi].data_label
                            pdl.has_text_frame = True; pdl.text_frame.text = str(txt)
                            # a per-point dLbl does not inherit the series dLblPos - set it again
                            if want: pdl.position = getattr(XL_LABEL_POSITION, want)
                            for para in pdl.text_frame.paragraphs:
                                for run in para.runs:
                                    run.font.size = Pt(half_pt(s.get("labelsize", 9))); run.font.name = FONT
                                    if s.get("labelbold"): run.font.bold = True
                                    if s.get("labelcolor"): run.font.color.rgb = RGBColor.from_string(s["labelcolor"])
                        except Exception: pass
        hl = ch.get("highlight")
        if hl:
            try:
                pt = chart.series[hl["series"]].points[hl["point"]]
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb = RGBColor.from_string(hl["color"])
            except Exception: pass
        if ch.get("valuemax"):
            # shared scale across separate charts (small multiples): pin 0..max so bar
            # heights compare honestly between charts instead of each auto-scaling
            try:
                chart.value_axis.minimum_scale = 0
                chart.value_axis.maximum_scale = float(ch["valuemax"])
            except Exception: pass
        if ch.get("valueaxis") == "hide":
            # hiding the axis leaves its gridlines behind, which reads as stray rules
            # across the plot - kill both unless gridlines are explicitly asked for
            try: chart.value_axis.visible = False
            except Exception: pass
        if ch.get("valueaxis") == "hide" or ch.get("gridlines") == "hide":
            try: chart.value_axis.has_major_gridlines = False
            except Exception: pass
            try: chart.value_axis.has_minor_gridlines = False
            except Exception: pass
        if ch.get("cataxis") == "hide":
            # small multi-card layouts often carry their own custom axis labels as
            # translated text boxes (e.g. single-letter month labels) - the chart's
            # own category axis must not also render, or the two collide/garble.
            try: chart.category_axis.visible = False
            except Exception: pass
        else:
            try:
                # data-catlabelsize (pt) for small multiples: at 11pt PowerPoint rotates the
                # labels of a narrow chart and they swallow the plot area
                chart.category_axis.tick_labels.font.size = Pt(half_pt(float(ch.get("catlabelsize") or 11)))
                # House default is non-bold, matching the HTML preview's category labels;
                # data-catlabelbold="true" opts a chart back into bold labels.
                chart.category_axis.tick_labels.font.bold = bool(ch.get("catlabelbold"))
            except Exception: pass
            if ch.get("valueaxis") == "hide" or ch.get("gridlines") == "hide":
                # The preview draws nothing but bars and labels in this mode - no axis
                # line either. Suppress the category axis's own line to match.
                try: chart.category_axis.format.line.fill.background()
                except Exception: pass
            rot = ch.get("catlabelRotation")
            if rot not in (None, ""):
                # Native chart category-axis labels rotate fine in PowerPoint (unlike
                # HTML div text, which the translator cannot rotate) - set it directly
                # via the axis's bodyPr, which tick_labels.font already created.
                try:
                    from pptx.oxml.ns import qn
                    catAx = chart.category_axis._element
                    txPr = catAx.find(qn('c:txPr'))
                    bodyPr = txPr.find(qn('a:bodyPr'))
                    if bodyPr is None:
                        bodyPr = txPr.makeelement(qn('a:bodyPr'), {})
                        txPr.insert(0, bodyPr)
                    bodyPr.set('rot', str(int(round(float(rot) * 60000))))
                except Exception as e:
                    print(f"  ! catlabelRotation failed: {e}")

def add_native_charts(pptx_path, charts):
    from pptx import Presentation as PPTXPres
    prs = PPTXPres(str(pptx_path))
    add_charts_to_slide(prs.slides[0], charts)
    prs.save(str(pptx_path))

def main():
    _cli.usage(__doc__, sys.argv[1:], 2, 2)
    geom, outp = Path(sys.argv[1]), Path(sys.argv[2])
    data = json.loads(geom.read_text(encoding="utf-8"))
    # Per-output work dir, tolerant of a stale lock (e.g. AV/handle): if it can't be
    # cleared, use a fresh suffixed dir rather than failing - so we never need to kill
    # a process (or the user's open PowerPoint) to release a lock.
    base = outp.parent / ("_build_" + re.sub(r"[^A-Za-z0-9_-]", "_", outp.stem))
    work = base
    if work.exists():
        try: shutil.rmtree(work)
        except OSError:
            n = 1
            while (outp.parent / f"{base.name}_{n}").exists(): n += 1
            work = outp.parent / f"{base.name}_{n}"
    work.mkdir(parents=True, exist_ok=True)
    unp = work/"unpacked"
    office_io.unpack(TPL, unp)
    _pf.patch(str(unp))
    (unp/"ppt"/"slides"/"_rels").mkdir(parents=True, exist_ok=True)
    idg = IdGen()
    media = MediaRegistry(unp)
    slide_rels = SlideRels()
    (unp/"ppt"/"slides"/"slide1.xml").write_text(build_slide_xml(data, idg, media, slide_rels), encoding="utf-8")
    media.ensure_content_types()
    (unp/"ppt"/"slides"/"_rels"/"slide1.xml.rels").write_text(
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="%s/slideLayout" Target="../slideLayouts/%s"/>%s</Relationships>'
        % (RT, layout_target(data, unp), slide_rels.rels_xml()),
        encoding="utf-8")
    register(unp)
    try:
        office_io.pack(unp, outp)
    except ValueError as e:
        print("PACK FAILED:", e); sys.exit(1)
    charts = data.get("charts") or []
    if charts:
        add_native_charts(outp, charts)
    print("shapes emitted:", idg.n - 1, "| native charts:", len(charts), "| ->", outp.name)

if __name__ == "__main__":
    main()
