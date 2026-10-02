"""Tests for image translation: <img> and non-gradient background-image become real
<p:pic> pictures in the .pptx.

- Pure-Python crop/letterbox maths and ellipse detection: no browser, no template.
- A hand-written geom.json round trip through html_to_pptx.py: no browser, needs the
  configured template (via scripts/config.py) like the rest of tests/test_translator.py.
- Extractor test (guarded with importorskip("playwright")): a real slide HTML with an
  <img>, a background-image:cover div, and a circle-cropped <img>, run through
  extract_geometry.py in a real browser.
"""
import json
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import html_to_pptx as H  # noqa: E402

from conftest import SKILL_ROOT, run_py  # noqa: E402


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _tiny_png(path, size=(40, 40), color=(200, 40, 40, 255)):
    from PIL import Image
    Image.new("RGBA", size, color).save(path)
    return path


# ---------------------------------------------------------------------------
# _cover_crop: "cover" (scale=None) and "none" (scale=1.0) srcRect maths
# ---------------------------------------------------------------------------

def test_cover_crop_wide_image_in_tall_box_centered():
    # natural 200x100 (2:1) into a 100x200 (1:2, tall) box: scale = max(0.5, 2) = 2,
    # so the image is 400x200 once scaled - all the excess is horizontal.
    l, t, r, b = H._cover_crop(200, 100, 100, 200, 0.5, 0.5)
    assert t == pytest.approx(0.0)
    assert b == pytest.approx(0.0)
    assert l == pytest.approx(0.375)
    assert r == pytest.approx(0.375)


def test_cover_crop_tall_image_in_wide_box_centered():
    # natural 100x200 (1:2) into a 200x100 (2:1, wide) box: scale = max(2, 0.5) = 2,
    # scaled to 200x400 - all the excess is vertical.
    l, t, r, b = H._cover_crop(100, 200, 200, 100, 0.5, 0.5)
    assert l == pytest.approx(0.0)
    assert r == pytest.approx(0.0)
    assert t == pytest.approx(0.375)
    assert b == pytest.approx(0.375)


def test_cover_crop_pos_0_puts_all_excess_on_the_trailing_edge():
    l, t, r, b = H._cover_crop(200, 100, 100, 200, 0.0, 0.5)
    assert l == pytest.approx(0.0)
    assert r == pytest.approx(0.75)


def test_cover_crop_pos_1_puts_all_excess_on_the_leading_edge():
    l, t, r, b = H._cover_crop(200, 100, 100, 200, 1.0, 0.5)
    assert l == pytest.approx(0.75)
    assert r == pytest.approx(0.0)


def test_cover_crop_none_fit_uses_fixed_scale_1():
    # "none": natural size (100x100) is larger than a 40x40 box on both axes; crop
    # (not scale) removes the excess, centred.
    l, t, r, b = H._cover_crop(100, 100, 40, 40, 0.5, 0.5, scale=1.0)
    assert l == pytest.approx(0.3)
    assert r == pytest.approx(0.3)
    assert t == pytest.approx(0.3)
    assert b == pytest.approx(0.3)


# ---------------------------------------------------------------------------
# _contain_box: letterbox maths for "contain" / "scale-down"
# ---------------------------------------------------------------------------

def test_contain_letterboxes_a_tall_image_into_a_square_box():
    # natural 100x200 into a 200x200 box: scale = min(2, 1) = 1 -> 100x200, centred
    # horizontally (50px each side), flush vertically.
    x, y, w, h = H._contain_box(100, 200, 0, 0, 200, 200, 0.5, 0.5)
    assert (x, y, w, h) == pytest.approx((50.0, 0.0, 100.0, 200.0))


def test_contain_respects_box_offset_and_position():
    x, y, w, h = H._contain_box(100, 200, 10, 10, 200, 200, 0.0, 1.0)
    assert x == pytest.approx(10.0)          # pos_x 0 -> flush left within the box
    assert y == pytest.approx(10.0 + 0.0)    # no vertical letterbox to place


def test_scale_down_never_upscales_a_smaller_image():
    # natural 50x50 into a 200x200 box: plain "contain" would scale up 4x; scale-down
    # clamps to 1x so a small image stays small, centred.
    x, y, w, h = H._contain_box(50, 50, 0, 0, 200, 200, 0.5, 0.5, only_shrink=True)
    assert (x, y, w, h) == pytest.approx((75.0, 75.0, 50.0, 50.0))


def test_contain_without_shrink_flag_upscales_freely():
    x, y, w, h = H._contain_box(50, 50, 0, 0, 200, 200, 0.5, 0.5, only_shrink=False)
    assert (w, h) == pytest.approx((200.0, 200.0))


# ---------------------------------------------------------------------------
# ellipse detection: radius >= half the shorter side -> a true circle geometry
# ---------------------------------------------------------------------------

def _image(**over):
    base = {"sel": "img.x", "x": 10, "y": 10, "w": 120, "h": 120, "src": None,
            "natural_w": 120, "natural_h": 120, "fit": "fill", "pos_x": 0.5, "pos_y": 0.5,
            "radius": 0, "opacity": 1, "z": 0}
    base.update(over)
    return base


def test_full_radius_square_image_becomes_an_ellipse(tmp_path):
    png = _tiny_png(tmp_path / "circle.png", size=(120, 120))
    idg = H.IdGen()
    media = H.MediaRegistry(tmp_path / "unpacked")
    rels = H.SlideRels()
    xml = H.pic_shape(idg, _image(src=str(png), radius=60), media, rels)
    assert 'prst="ellipse"' in xml
    assert 'prst="rect"' not in xml


def test_small_radius_square_image_stays_a_rect(tmp_path):
    png = _tiny_png(tmp_path / "rounded.png", size=(120, 120))
    idg = H.IdGen()
    media = H.MediaRegistry(tmp_path / "unpacked")
    rels = H.SlideRels()
    xml = H.pic_shape(idg, _image(src=str(png), radius=8), media, rels)
    assert 'prst="rect"' in xml
    assert 'prst="ellipse"' not in xml


def test_opacity_below_1_emits_alphaModFix(tmp_path):
    png = _tiny_png(tmp_path / "img.png")
    idg = H.IdGen()
    media = H.MediaRegistry(tmp_path / "unpacked")
    rels = H.SlideRels()
    xml = H.pic_shape(idg, _image(src=str(png), opacity=0.4), media, rels)
    assert '<a:alphaModFix amt="40000"/>' in xml


def test_full_opacity_omits_alphaModFix(tmp_path):
    png = _tiny_png(tmp_path / "img.png")
    idg = H.IdGen()
    media = H.MediaRegistry(tmp_path / "unpacked")
    rels = H.SlideRels()
    xml = H.pic_shape(idg, _image(src=str(png), opacity=1.0), media, rels)
    assert "alphaModFix" not in xml


def test_picture_is_named_picture_n_not_pic_n(tmp_path):
    """The selection pane should read 'Picture 2', not the old internal 'pic2' - see
    the naming fix in html_to_pptx.rect_shape/_txsp/pic_shape."""
    png = _tiny_png(tmp_path / "img.png")
    idg = H.IdGen()
    media = H.MediaRegistry(tmp_path / "unpacked")
    rels = H.SlideRels()
    xml = H.pic_shape(idg, _image(src=str(png)), media, rels)
    assert 'name="Picture 2"' in xml
    assert 'name="pic2"' not in xml


# ---------------------------------------------------------------------------
# MediaRegistry.ensure_content_types: adds a missing Default extension, skips one
# already declared (the plain template may already declare some of these).
# ---------------------------------------------------------------------------

def test_ensure_content_types_adds_missing_extension_once(tmp_path):
    unp = tmp_path / "unpacked"
    (unp / "ppt" / "media").mkdir(parents=True)
    ct = unp / "[Content_Types].xml"
    ct.write_text(
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/></Types>',
        encoding="utf-8",
    )
    media = H.MediaRegistry(unp)
    media.exts_used.add("webp")
    media.ensure_content_types()
    s = ct.read_text(encoding="utf-8")
    assert s.count('Extension="webp"') == 1
    assert 'ContentType="image/webp"' in s
    # calling it again must not duplicate the Default
    media.ensure_content_types()
    assert ct.read_text(encoding="utf-8").count('Extension="webp"') == 1


# ---------------------------------------------------------------------------
# Round trip without a browser: rect + image + text (z-interleaved) through the
# real single-slide translator CLI, then opened with python-pptx.
# ---------------------------------------------------------------------------

def test_translator_round_trip_places_picture_dedupes_media_and_preserves_z_order(tmp_path):
    png = _tiny_png(tmp_path / "photo.png", size=(40, 40))
    geom = {
        "w": 1280, "h": 720, "layout": None, "bodyBg": None,
        "rects": [{"x": 40, "y": 40, "w": 300, "h": 200, "fill": "rgb(20,20,20)",
                   "border": None, "radius": 0, "z": 0}],
        "texts": [{"x": 40, "y": 600, "w": 400, "h": 50, "lines": 1, "lhpx": 20,
                   "align": "left", "anchor": "t", "pl": 0, "pr": 0, "pt": 0, "pb": 0,
                   "paras": [[{"t": "Caption", "size": 14, "weight": 400,
                               "color": "rgb(255,255,255)", "italic": False,
                               "upper": False, "ls": 0}]], "z": 2}],
        # two images, SAME src file, on either side of the text in z-order -
        # exercises both z-interleaving and media dedup in one build.
        "images": [
            {"sel": "img#a", "x": 100, "y": 100, "w": 200, "h": 150, "src": str(png),
             "natural_w": 40, "natural_h": 40, "fit": "fill", "pos_x": 0.5, "pos_y": 0.5,
             "radius": 0, "opacity": 1, "z": 1},
            {"sel": "img#b", "x": 900, "y": 50, "w": 80, "h": 80, "src": str(png),
             "natural_w": 40, "natural_h": 40, "fit": "fill", "pos_x": 0.5, "pos_y": 0.5,
             "radius": 0, "opacity": 1, "z": 3},
        ],
        "charts": [], "warnings": [],
    }
    geom_path = tmp_path / "geom.json"
    geom_path.write_text(json.dumps(geom), encoding="utf-8")
    out = tmp_path / "out.pptx"

    r = run_py([str(SKILL_ROOT / "scripts" / "html_to_pptx.py"), str(geom_path), str(out)],
               cwd=SKILL_ROOT)
    assert r.returncode == 0, r.stdout + r.stderr
    assert out.exists()

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(out))
    slide = prs.slides[0]
    shapes = list(slide.shapes)

    types = [s.shape_type for s in shapes]
    assert types == [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE,
                      MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PICTURE], \
        "shape order in spTree must follow z: rect, image, text, image"

    pic1 = shapes[1]
    assert abs(pic1.left - 100 * 9525) <= 1
    assert abs(pic1.top - 100 * 9525) <= 1
    assert abs(pic1.width - 200 * 9525) <= 1
    assert abs(pic1.height - 150 * 9525) <= 1

    with zipfile.ZipFile(H.TPL) as tz:
        baseline = {n for n in tz.namelist() if n.startswith("ppt/media/")}
    with zipfile.ZipFile(out) as z:
        media_parts = [n for n in z.namelist() if n.startswith("ppt/media/")]
        added = [n for n in media_parts if n not in baseline]
        assert len(added) == 1, \
            f"the same source file must only be embedded once, found new parts {added}"
        ct = z.read("[Content_Types].xml").decode("utf-8")
        assert 'Extension="png"' in ct


# ---------------------------------------------------------------------------
# Extractor test: a real slide HTML with an <img>, a background-image:cover div,
# and a circle-cropped <img>, run through extract_geometry.py in a real browser.
# ---------------------------------------------------------------------------

pytestmark_browser = pytest.mark.browser


@pytest.mark.browser
def test_extract_geometry_captures_img_and_background_image(tmp_path):
    pytest.importorskip("playwright")
    photo = _tiny_png(tmp_path / "photo.png", size=(200, 100), color=(30, 60, 200, 255))
    bg = _tiny_png(tmp_path / "bg.png", size=(300, 150), color=(10, 200, 60, 255))
    headshot = _tiny_png(tmp_path / "headshot.png", size=(120, 120), color=(200, 200, 10, 255))

    html = tmp_path / "slide.html"
    html.write_text(
        "<!DOCTYPE html><html><head><meta charset='UTF-8'></head>"
        "<body style=\"margin:0;padding:0;width:1280px;height:720px\">"
        "<div class=\"slide\" style=\"width:1280px;height:720px;position:relative;"
        "background:#fff;overflow:hidden\">"
        "<img src=\"photo.png\" style=\"position:absolute;left:20px;top:20px;"
        "width:200px;height:100px;object-fit:contain\">"
        "<div style=\"position:absolute;left:260px;top:20px;width:300px;height:150px;"
        "background-image:url('bg.png');background-size:cover;background-position:50% 50%\">"
        "</div>"
        "<img src=\"headshot.png\" style=\"position:absolute;left:600px;top:20px;"
        "width:120px;height:120px;border-radius:50%;object-fit:cover\">"
        "</div></body></html>",
        encoding="utf-8",
    )
    out = tmp_path / "geom.json"
    r = run_py([str(SKILL_ROOT / "scripts" / "extract_geometry.py"), str(html), str(out)],
               cwd=SKILL_ROOT, timeout=60)
    assert r.returncode == 0, r.stdout + r.stderr
    data = json.loads(out.read_text(encoding="utf-8"))

    images = data.get("images") or []
    assert len(images) == 3, [im["sel"] for im in images]

    by_sel = {im["sel"]: im for im in images}
    img_entry = next(im for im in images if im["src"].endswith("photo.png"))
    assert img_entry["fit"] == "contain"
    assert abs(img_entry["w"] - 200) < 1 and abs(img_entry["h"] - 100) < 1

    bg_entry = next(im for im in images if im["src"].endswith("bg.png"))
    assert bg_entry["fit"] == "cover"
    assert abs(bg_entry["pos_x"] - 0.5) < 0.01
    assert abs(bg_entry["pos_y"] - 0.5) < 0.01

    circle_entry = next(im for im in images if im["src"].endswith("headshot.png"))
    assert circle_entry["fit"] == "cover"
    assert circle_entry["radius"] >= min(circle_entry["w"], circle_entry["h"]) / 2 - 0.5

    for w in data.get("warnings") or []:
        assert "img" not in w["why"].lower() and "image" not in w["why"].lower(), w
