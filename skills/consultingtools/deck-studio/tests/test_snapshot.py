"""Tests for scripts/snapshot.py.

- Manifest allocation and record construction against a fake MANIFEST.json copied into a tmp
  dir (ASSET_BANK is monkeypatched to that tmp dir - the real asset bank is never touched):
  next id, slug, dedupe by sha256, tags merge.
- Slide-title extraction: from a .pptx built with python-pptx (title placeholder, and the
  no-title-placeholder fallback), and from an HTML string (class="title", <title>, default).
- --help / missing-file CLI contract, via subprocess like the rest of the suite.

No PowerPoint and no browser needed - render_pptx_slide / render_pptx_all / render_html_slide
are not exercised here.
"""
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import snapshot as S  # noqa: E402

from conftest import SKILL_ROOT, run_py  # noqa: E402


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _tiny_png(path, size=(40, 40), color=(200, 40, 40, 255)):
    from PIL import Image
    Image.new("RGBA", size, color).save(path)
    return path


def _fake_bank(tmp_path, assets):
    """A tmp asset-bank dir with just enough MANIFEST.json to exercise load/save/file_in_bank."""
    bank = tmp_path / "asset_bank"
    bank.mkdir()
    (bank / "MANIFEST.json").write_text(
        json.dumps({"created": "2026-09-18", "assets": assets}, indent=1, ensure_ascii=False),
        encoding="utf-8",
    )
    return bank


# ---------------------------------------------------------------------------
# next_id / slugify
# ---------------------------------------------------------------------------

def test_next_id_increments_past_highest_existing():
    assets = [{"id": "a0001"}, {"id": "a0526"}, {"id": "a0003"}]
    assert S.next_id(assets) == "a0527"


def test_next_id_with_no_assets_starts_at_a0001():
    assert S.next_id([]) == "a0001"


def test_next_id_ignores_malformed_ids():
    assets = [{"id": "a0010"}, {"id": "not-an-id"}, {}]
    assert S.next_id(assets) == "a0011"


def test_slugify_lowercases_and_replaces_punctuation():
    assert S.slugify("The Q3 Result, Presented!") == "the_q3_result_presented"


def test_slugify_truncates_to_maxlen_without_trailing_underscore():
    s = S.slugify("word " * 30, maxlen=20)
    assert len(s) <= 20
    assert not s.endswith("_")


def test_slugify_empty_text_falls_back_to_capture():
    assert S.slugify("") == "capture"
    assert S.slugify("!!!") == "capture"


# ---------------------------------------------------------------------------
# build_record
# ---------------------------------------------------------------------------

def test_build_record_fields_for_a_pptx_source(tmp_path):
    png = _tiny_png(tmp_path / "cap.png", size=(64, 32))
    rec = S.build_record(
        "a0527", "screenshot/a0527_x.png", "mydeck", r"C:\decks\mydeck.pptx", 3,
        "The Q3 result", None, ["demo"], png, "2026-09-19T12:00:00+00:00",
    )
    assert rec["id"] == "a0527"
    assert rec["file"] == "screenshot/a0527_x.png"
    assert rec["category"] == "screenshot"
    assert rec["ext"] == "png"
    assert rec["bytes"] == png.stat().st_size
    assert rec["dims"] == [64, 32]
    assert rec["alpha"] is False
    assert rec["sources"] == [r"C:\decks\mydeck.pptx slide 3"]
    assert rec["source_slides"] == []
    assert rec["slide_titles"] == ["The Q3 result"]
    assert rec["pptx_names"] == []
    assert rec["descr"] == []
    assert rec["description"] == "Slide capture: The Q3 result (mydeck, slide 3)"
    assert rec["tags"] == ["snapshot", "slide capture", "mydeck", "demo"]
    assert len(rec["sha256"]) == 64
    assert rec["snapshot"] == {
        "source": r"C:\decks\mydeck.pptx", "slide": 3, "taken": "2026-09-19T12:00:00+00:00",
    }


def test_build_record_html_source_has_no_slide_number_in_sources_or_description(tmp_path):
    png = _tiny_png(tmp_path / "cap.png")
    rec = S.build_record(
        "a0001", "screenshot/a0001_x.png", "exhibit", r"C:\slides\exhibit.html", None,
        "Some title", None, [], png, "t",
    )
    assert rec["sources"] == [r"C:\slides\exhibit.html"]
    assert rec["description"] == "Slide capture: Some title (exhibit)"
    assert rec["snapshot"]["slide"] is None


def test_build_record_explicit_desc_overrides_default(tmp_path):
    png = _tiny_png(tmp_path / "cap.png")
    rec = S.build_record("a0001", "screenshot/a0001_x.png", "deck", "src", 1, "Title",
                          "My custom description", [], png, "t")
    assert rec["description"] == "My custom description"


def test_build_record_tags_merge_dedupes_base_tags():
    png_path = None
    # base tags are ["snapshot", "slide capture", <deck stem>] - passing one again must not
    # duplicate it, but a genuinely new tag is appended.
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        png = _tiny_png(Path(d) / "cap.png")
        rec = S.build_record("a0001", "screenshot/a0001_x.png", "deck", "src", 1, "Title",
                              None, ["snapshot", "hero"], png, "t")
        assert rec["tags"] == ["snapshot", "slide capture", "deck", "hero"]


# ---------------------------------------------------------------------------
# file_in_bank: dedupe by sha256, id allocation, manifest write
# ---------------------------------------------------------------------------

def test_file_in_bank_creates_new_record_and_writes_manifest(tmp_path, monkeypatch):
    bank = _fake_bank(tmp_path, [{"id": "a0001", "sha256": "0" * 64}])
    monkeypatch.setattr(S, "ASSET_BANK", bank)

    png = _tiny_png(tmp_path / "src.png", color=(10, 20, 30, 255))
    record, created = S.file_in_bank(png, "src.pptx", 5, "A Title", None, [], "deck")

    assert created is True
    assert record["id"] == "a0002"
    dest = bank / record["file"]
    assert dest.exists()
    assert dest.read_bytes() == png.read_bytes()

    manifest = json.loads((bank / "MANIFEST.json").read_text(encoding="utf-8"))
    assert [a["id"] for a in manifest["assets"]] == ["a0001", "a0002"]


def test_file_in_bank_dedupes_identical_image_by_sha256(tmp_path, monkeypatch):
    bank = _fake_bank(tmp_path, [])
    monkeypatch.setattr(S, "ASSET_BANK", bank)

    png = _tiny_png(tmp_path / "src.png", color=(50, 60, 70, 255))
    rec1, created1 = S.file_in_bank(png, "src.pptx", 1, "First capture", None, [], "deck")
    assert created1 is True

    # a second, byte-identical PNG (different path, different slide number/title) must not
    # add a second record - only the pixels are deduped on.
    png2 = _tiny_png(tmp_path / "src2.png", color=(50, 60, 70, 255))
    rec2, created2 = S.file_in_bank(png2, "src.pptx", 2, "Second capture", None, [], "deck")

    assert created2 is False
    assert rec2["id"] == rec1["id"]

    manifest = json.loads((bank / "MANIFEST.json").read_text(encoding="utf-8"))
    assert len(manifest["assets"]) == 1
    assert len(list((bank / "screenshot").glob("*.png"))) == 1


def test_file_in_bank_different_images_get_different_ids(tmp_path, monkeypatch):
    bank = _fake_bank(tmp_path, [])
    monkeypatch.setattr(S, "ASSET_BANK", bank)

    png_a = _tiny_png(tmp_path / "a.png", color=(1, 2, 3, 255))
    png_b = _tiny_png(tmp_path / "b.png", color=(9, 8, 7, 255))
    rec_a, _ = S.file_in_bank(png_a, "src.pptx", 1, "Slide A", None, [], "deck")
    rec_b, _ = S.file_in_bank(png_b, "src.pptx", 2, "Slide B", None, [], "deck")

    assert rec_a["id"] != rec_b["id"]
    manifest = json.loads((bank / "MANIFEST.json").read_text(encoding="utf-8"))
    assert len(manifest["assets"]) == 2


# ---------------------------------------------------------------------------
# slide_title_from_pptx
# ---------------------------------------------------------------------------

def _one_slide_pptx_with_title(path, title_text):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide layout
    slide.shapes.title.text = title_text
    prs.save(str(path))


def _one_slide_pptx_without_title_placeholder(path, body_text):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout: no placeholders
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = body_text
    prs.save(str(path))


def test_slide_title_from_pptx_reads_title_placeholder(tmp_path):
    p = tmp_path / "deck.pptx"
    _one_slide_pptx_with_title(p, "We are a consultancy collective")
    assert S.slide_title_from_pptx(p, 1) == "We are a consultancy collective"


def test_slide_title_from_pptx_falls_back_to_first_text_frame(tmp_path):
    p = tmp_path / "deck.pptx"
    _one_slide_pptx_without_title_placeholder(p, "Only a body text box")
    assert S.slide_title_from_pptx(p, 1) == "Only a body text box"


def test_slide_title_from_pptx_out_of_range_raises(tmp_path):
    p = tmp_path / "deck.pptx"
    _one_slide_pptx_with_title(p, "Slide one")
    with pytest.raises(ValueError):
        S.slide_title_from_pptx(p, 2)


# ---------------------------------------------------------------------------
# title_from_html
# ---------------------------------------------------------------------------

def test_title_from_html_reads_class_title_element():
    html = '<div class="slide"><div class="title">The Q3 result was strong</div></div>'
    assert S.title_from_html(html) == "The Q3 result was strong"


def test_title_from_html_strips_nested_tags_and_entities():
    html = '<h1 class="title">Revenue &amp; <b>EBITDA</b> grew</h1>'
    assert S.title_from_html(html) == "Revenue & EBITDA grew"


def test_title_from_html_falls_back_to_title_tag():
    html = "<head><title>Fallback Deck Title</title></head><body></body>"
    assert S.title_from_html(html) == "Fallback Deck Title"


def test_title_from_html_prefers_class_title_over_title_tag():
    html = '<head><title>Doc Title</title></head><body><div class="title">On-slide title</div></body>'
    assert S.title_from_html(html) == "On-slide title"


def test_title_from_html_defaults_to_untitled_when_neither_present():
    assert S.title_from_html("<div>no title anywhere</div>") == "untitled"


# ---------------------------------------------------------------------------
# CLI contract: --help / -h exit 2, missing file exits 1
# (also exercised generically by tests/test_cli.py's discovery of every scripts/*.py)
# ---------------------------------------------------------------------------

def test_help_prints_usage_and_exits_2():
    r = run_py([str(SKILL_ROOT / "scripts" / "snapshot.py"), "--help"], cwd=SKILL_ROOT)
    assert r.returncode == 2
    assert "Usage:" in r.stdout


def test_missing_source_file_exits_1():
    r = run_py([str(SKILL_ROOT / "scripts" / "snapshot.py"), "does_not_exist.pptx", "1"],
               cwd=SKILL_ROOT)
    assert r.returncode == 1
    assert "FAILED" in r.stdout


def test_unsupported_file_type_exits_1(tmp_path):
    bogus = tmp_path / "deck.txt"
    bogus.write_text("not a deck", encoding="utf-8")
    r = run_py([str(SKILL_ROOT / "scripts" / "snapshot.py"), str(bogus)], cwd=SKILL_ROOT)
    assert r.returncode == 1
    assert "FAILED" in r.stdout


def test_pptx_without_slide_number_or_all_exits_1(tmp_path):
    p = tmp_path / "deck.pptx"
    _one_slide_pptx_with_title(p, "Slide one")
    r = run_py([str(SKILL_ROOT / "scripts" / "snapshot.py"), str(p)], cwd=SKILL_ROOT)
    assert r.returncode == 1
    assert "FAILED" in r.stdout
