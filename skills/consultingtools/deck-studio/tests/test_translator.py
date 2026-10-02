"""Unit tests for scripts/html_to_pptx.py's pure-Python translation logic.

These build hand-written geometry dicts shaped like scripts/extract_geometry.py's JS
output (rects/texts/charts, see its docstring and JS block) and feed them straight into
html_to_pptx functions - no browser, no Playwright, no template unpacking needed for most
of these (see test_layouts.py for the template-dependent layout_target tests).
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import html_to_pptx as H  # noqa: E402


def _run(t="hi", **over):
    run = {"t": t, "size": 14, "weight": 400, "color": "rgb(0,0,0)",
           "italic": False, "upper": False, "ls": 0}
    run.update(over)
    return run


def _text(**over):
    base = {"x": 10, "y": 10, "w": 300, "h": 40, "lines": 1, "lhpx": 16.8,
            "align": "left", "anchor": "t", "pl": 0, "pr": 0, "pt": 0, "pb": 0,
            "paras": [[_run()]]}
    base.update(over)
    return base


# ---------------------------------------------------------------------------
# IdGen
# ---------------------------------------------------------------------------

def test_idgen_ids_unique_across_500_shapes():
    idg = H.IdGen()
    ids = [idg() for _ in range(500)]
    assert len(ids) == len(set(ids)), "IdGen produced duplicate shape ids"
    assert all(isinstance(i, int) for i in ids)


# ---------------------------------------------------------------------------
# wrap="none" vs wrap="square"
# ---------------------------------------------------------------------------

def test_single_line_short_text_gets_wrap_none():
    idg = H.IdGen()
    xml = H.text_shape(idg, _text(paras=[[_run("Short label")]], lines=1))
    assert 'wrap="none"' in xml
    assert 'wrap="square"' not in xml


def test_same_text_with_two_lines_gets_wrap_square():
    idg = H.IdGen()
    xml = H.text_shape(idg, _text(paras=[[_run("Short label")]], lines=2))
    assert 'wrap="square"' in xml
    assert 'wrap="none"' not in xml


def test_long_text_over_34_chars_never_gets_wrap_none_even_at_one_line():
    idg = H.IdGen()
    long_text = "A" * 40
    xml = H.text_shape(idg, _text(paras=[[_run(long_text)]], lines=1))
    assert 'wrap="square"' in xml


# ---------------------------------------------------------------------------
# grouped (.pgroup) line spacing: spcPts (lhpx) vs spcPct (lhpct)
# ---------------------------------------------------------------------------

def _grouped(gparas):
    return {"x": 0, "y": 0, "w": 300, "h": 100, "align": "left", "anchor": "t",
            "pl": 10, "pr": 10, "pt": 5, "pb": 5, "grouped": True, "gparas": gparas}


def test_grouped_text_with_lhpx_emits_spcPts():
    idg = H.IdGen()
    gp = {"runs": [_run("Item one")], "align": "left", "spcBef": 0, "lhpx": 18.0}
    xml = H.text_shape(idg, _grouped([gp]))
    assert "<a:lnSpc><a:spcPts" in xml
    assert "<a:lnSpc><a:spcPct" not in xml


def test_grouped_text_with_lhpct_only_emits_spcPct():
    idg = H.IdGen()
    gp = {"runs": [_run("Item one")], "align": "left", "spcBef": 0, "lhpct": 120000}
    xml = H.text_shape(idg, _grouped([gp]))
    assert "<a:lnSpc><a:spcPct" in xml
    assert "<a:lnSpc><a:spcPts" not in xml


def test_grouped_multi_paragraph_spcBef_only_above_threshold():
    idg = H.IdGen()
    gparas = [
        {"runs": [_run("First")], "align": "left", "spcBef": 0, "lhpx": 18.0},
        {"runs": [_run("Second")], "align": "left", "spcBef": 6.0, "lhpx": 18.0},
    ]
    xml = H.text_shape(idg, _grouped(gparas))
    assert xml.count("<a:spcBef>") == 1  # only the second paragraph clears the 0.3pt-ish floor


# ---------------------------------------------------------------------------
# stacked-chart label position clamping (illegal labelpos -> inside_end)
# ---------------------------------------------------------------------------

def test_stacked_chart_illegal_labelpos_is_clamped_to_legal_value(capsys):
    from pptx import Presentation
    from pptx.enum.chart import XL_LABEL_POSITION

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    charts = [{
        "x": 1, "y": 1, "w": 5, "h": 4, "type": "stacked-column",
        "categories": ["A", "B"],
        "series": [{"name": "S1", "values": [1, 2], "color": "1E1E1E",
                    "labels": True, "labelpos": "outside_end"}],
        "legend": "top", "valueaxis": "show",
    }]
    H.add_charts_to_slide(slide, charts)
    chart = slide.shapes[0].chart
    pos = chart.series[0].data_labels.position
    assert pos == XL_LABEL_POSITION.INSIDE_END
    assert "illegal on a stacked chart" in capsys.readouterr().out


@pytest.mark.parametrize("legal_pos", ["center", "inside_end", "inside_base"])
def test_stacked_chart_legal_labelpos_passes_through_unclamped(legal_pos, capsys):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    charts = [{
        "x": 1, "y": 1, "w": 5, "h": 4, "type": "stacked-bar",
        "categories": ["A", "B"],
        "series": [{"name": "S1", "values": [1, 2], "color": "1E1E1E",
                    "labels": True, "labelpos": legal_pos}],
        "legend": "none", "valueaxis": "show",
    }]
    H.add_charts_to_slide(slide, charts)
    assert "illegal on a stacked chart" not in capsys.readouterr().out


def test_non_stacked_chart_outside_end_is_not_clamped(capsys):
    from pptx import Presentation
    from pptx.enum.chart import XL_LABEL_POSITION

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    charts = [{
        "x": 1, "y": 1, "w": 5, "h": 4, "type": "column",
        "categories": ["A", "B"],
        "series": [{"name": "S1", "values": [1, 2], "color": "1E1E1E",
                    "labels": True, "labelpos": "outside_end"}],
        "legend": "none", "valueaxis": "show",
    }]
    H.add_charts_to_slide(slide, charts)
    chart = slide.shapes[0].chart
    assert chart.series[0].data_labels.position == XL_LABEL_POSITION.OUTSIDE_END
    assert "illegal on a stacked chart" not in capsys.readouterr().out


# ---------------------------------------------------------------------------
# background promotion: an opaque full-bleed bottom rect becomes <p:bg>, not a shape
# ---------------------------------------------------------------------------

def test_full_bleed_bottom_rect_promoted_to_slide_bg_not_a_shape():
    idg = H.IdGen()
    data = {
        "bodyBg": None,
        "rects": [{"x": 0, "y": 0, "w": 1280, "h": 720, "fill": "rgb(238,238,238)",
                   "border": None, "radius": 0}],
        "texts": [],
    }
    xml = H.build_slide_xml(data, idg)
    assert "<p:bg>" in xml
    assert xml.count("<p:sp>") == 0
    assert "EEEEEE" in xml


def test_full_bleed_rect_with_border_is_not_promoted():
    """A bordered full-bleed rect is a deliberate frame/scrim, not a fake canvas colour."""
    idg = H.IdGen()
    data = {
        "bodyBg": None,
        "rects": [{"x": 0, "y": 0, "w": 1280, "h": 720, "fill": "rgb(238,238,238)",
                   "border": {"t": 2, "l": 2, "r": 2, "b": 2, "colT": "rgb(0,0,0)",
                              "colL": "rgb(0,0,0)", "colR": "rgb(0,0,0)", "colB": "rgb(0,0,0)"},
                   "radius": 0}],
        "texts": [],
    }
    xml = H.build_slide_xml(data, idg)
    assert "<p:bg>" not in xml
    assert xml.count("<p:sp>") == 1


def test_second_full_bleed_rect_stays_a_shape_scrim():
    """Only the bottom-of-stack full-bleed rect is promoted; a later one (an overlay
    scrim) is a deliberate real object and must stay a shape."""
    idg = H.IdGen()
    data = {
        "bodyBg": None,
        "rects": [
            {"x": 0, "y": 0, "w": 1280, "h": 720, "fill": "rgb(255,255,255)",
             "border": None, "radius": 0},
            {"x": 0, "y": 0, "w": 1280, "h": 720, "fill": "rgba(0,0,0,0.4)",
             "border": None, "radius": 0},
        ],
        "texts": [],
    }
    xml = H.build_slide_xml(data, idg)
    assert "<p:bg>" in xml
    assert xml.count("<p:sp>") == 1  # the rgba scrim remains a real shape


# ---------------------------------------------------------------------------
# colour parsing: rgb(), rgba() with alpha, #hex, transparent
# ---------------------------------------------------------------------------

def test_color_parses_rgb():
    assert H.color("rgb(255, 0, 0)") == ("FF0000", None)


def test_color_parses_rgba_with_alpha():
    hexv, alpha = H.color("rgba(255, 0, 0, 0.5)")
    assert hexv == "FF0000"
    assert alpha == 50000  # 0..100000 scale


def test_color_rgba_full_opacity_reports_no_alpha():
    hexv, alpha = H.color("rgba(10, 20, 30, 1)")
    assert hexv == "0A141E"
    assert alpha is None


def test_color_transparent_and_empty_are_none():
    assert H.color("transparent") == (None, None)
    assert H.color("") == (None, None)
    assert H.color(None) == (None, None)


def test_color_parses_hex():
    assert H.color("#FF00FF") == ("FF00FF", None)


# ---------------------------------------------------------------------------
# thin filled rects -> straight-connector line shapes
# ---------------------------------------------------------------------------

def _rect(**over):
    base = {"x": 20, "y": 30, "w": 200, "h": 6, "fill": "rgb(30,30,30)", "grad": "",
            "border": None, "radius": 0}
    base.update(over)
    return base


def test_thin_horizontal_rule_becomes_a_straight_connector():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(x=20, y=30, w=200, h=6))
    assert "<p:cxnSp>" in xml
    assert "<p:sp>" not in xml
    assert 'prst="line"' in xml
    assert 'name="Straight Connector 2"' in xml
    # centred on the box's vertical midline (y + h/2 = 33px -> 33*9525 EMU), full width,
    # stroke weight = the box's thickness (6px -> 6*9525 EMU), flat cap.
    assert '<a:off x="%d" y="%d"/>' % (H.px(20), H.px(33)) in xml
    assert '<a:ext cx="%d" cy="0"/>' % H.px(200) in xml
    assert '<a:ln w="%d" cap="flat">' % H.px(6) in xml
    assert "1E1E1E" in xml


def test_thin_vertical_rule_becomes_a_vertical_connector():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(x=50, y=10, w=4, h=150))
    assert "<p:cxnSp>" in xml
    # centred on the box's horizontal midline (x + w/2 = 52px), full height, cx=0.
    assert '<a:off x="%d" y="%d"/>' % (H.px(52), H.px(10)) in xml
    assert '<a:ext cx="0" cy="%d"/>' % H.px(150) in xml
    assert '<a:ln w="%d" cap="flat">' % H.px(4) in xml


def test_small_square_dot_is_not_treated_as_a_rule():
    """w == h (a tiny square icon, not a rule) stays an ordinary rect shape."""
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(x=0, y=0, w=6, h=6))
    assert "<p:cxnSp>" not in xml
    assert 'prst="rect"' in xml


def test_thin_rect_with_border_stays_a_rect_not_a_line():
    idg = H.IdGen()
    b = {"t": 1, "l": 1, "r": 1, "b": 1, "colT": "rgb(0,0,0)", "colL": "rgb(0,0,0)",
         "colR": "rgb(0,0,0)", "colB": "rgb(0,0,0)"}
    xml = H.rect_shape(idg, _rect(border=b))
    assert "<p:cxnSp>" not in xml
    assert 'prst="rect"' in xml


def test_thin_rect_with_gradient_stays_a_rect_not_a_line():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(fill="", grad="linear-gradient(rgb(0,0,0), rgb(255,255,255))"))
    assert "<p:cxnSp>" not in xml


def test_thin_rect_with_radius_stays_a_rect_not_a_line():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(radius=3))
    assert "<p:cxnSp>" not in xml


def test_line_shapes_keep_z_order_with_rects_and_text():
    """A rule mixed in with an ordinary rect and a text box must still land in the
    same document order in build_slide_xml (z-order is preserved regardless of which
    shapes end up as <p:sp> vs <p:cxnSp>)."""
    idg = H.IdGen()
    data = {
        "bodyBg": None,
        "rects": [
            _rect(x=0, y=0, w=100, h=100, fill="rgb(200,200,200)", z=0),   # ordinary rect
            _rect(x=0, y=50, w=200, h=6, z=1),                              # rule -> line
        ],
        "texts": [dict(_text(), z=2)],
    }
    xml = H.build_slide_xml(data, idg)
    order = [xml.index(tag) for tag in ("<p:sp>", "<p:cxnSp>") if tag in xml]
    # first <p:sp> (the ordinary rect) appears before <p:cxnSp> (the rule)
    assert xml.index("<p:sp>") < xml.index("<p:cxnSp>")


# ---------------------------------------------------------------------------
# shape naming: the selection pane should show names PowerPoint users recognise
# ---------------------------------------------------------------------------

def test_rect_is_named_rectangle_n():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(x=0, y=0, w=100, h=50, radius=0))
    assert 'name="Rectangle 2"' in xml


def test_rounded_rect_is_named_rounded_rectangle_n():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(x=0, y=0, w=100, h=50, radius=8))
    assert 'name="Rounded Rectangle 2"' in xml


def test_circle_rect_is_named_oval_n():
    idg = H.IdGen()
    xml = H.rect_shape(idg, _rect(x=0, y=0, w=60, h=60, radius=30))
    assert 'name="Oval 2"' in xml


def test_text_box_is_named_textbox_n():
    idg = H.IdGen()
    xml = H.text_shape(idg, _text())
    assert 'name="TextBox 2"' in xml


# ---------------------------------------------------------------------------
# horizontal bar charts: category order, label weight, axis line
# ---------------------------------------------------------------------------

def _bar_chart(**over):
    base = {
        "x": 1, "y": 1, "w": 6, "h": 4, "type": "bar",
        "categories": ["Logo", "Beta", "Gamma", "Delta", "Epsilon"],
        "series": [{"name": "S1", "values": [5, 4, 3, 2, 1], "color": "1E1E1E"}],
        "legend": "none", "valueaxis": "show",
    }
    base.update(over)
    return base


def test_horizontal_bar_chart_reverses_category_axis_so_first_category_is_on_top():
    from pptx import Presentation
    from pptx.enum.chart import XL_AXIS_CROSSES

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart()])
    chart = slide.shapes[0].chart
    assert chart.category_axis.reverse_order is True
    # Rendered and visually confirmed in PowerPoint: pinning value_axis.crosses to
    # MAXIMUM here (an approach that looks plausible on paper) actually reverses the
    # value scale and drops the bar fill - reverse_order alone is correct, so the
    # value axis crossing must be left at its untouched default (autoZero).
    assert chart.value_axis.crosses == XL_AXIS_CROSSES.AUTOMATIC


def test_vertical_column_chart_is_not_reversed():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart(type="column")])
    chart = slide.shapes[0].chart
    assert chart.category_axis.reverse_order is False


def test_stacked_bar_chart_is_also_reversed():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart(type="stacked-bar")])
    chart = slide.shapes[0].chart
    assert chart.category_axis.reverse_order is True


def test_bar_chart_category_labels_are_not_bold_by_default():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart()])
    chart = slide.shapes[0].chart
    assert chart.category_axis.tick_labels.font.bold is False


def test_bar_chart_category_labels_bold_when_html_asks_for_it():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart(catlabelbold=True)])
    chart = slide.shapes[0].chart
    assert chart.category_axis.tick_labels.font.bold is True


def test_bar_chart_category_axis_line_hidden_when_valueaxis_hidden():
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart(valueaxis="hide")])
    chart = slide.shapes[0].chart
    catAx = chart.category_axis._element
    spPr = catAx.find(qn('c:spPr'))
    assert spPr is not None
    assert spPr.find(qn('a:ln')).find(qn('a:noFill')) is not None


def test_bar_chart_category_axis_line_hidden_when_gridlines_hidden():
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart(gridlines="hide")])
    chart = slide.shapes[0].chart
    catAx = chart.category_axis._element
    spPr = catAx.find(qn('c:spPr'))
    assert spPr is not None
    assert spPr.find(qn('a:ln')).find(qn('a:noFill')) is not None


def test_bar_chart_category_axis_line_shown_by_default():
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    H.add_charts_to_slide(slide, [_bar_chart()])
    chart = slide.shapes[0].chart
    catAx = chart.category_axis._element
    spPr = catAx.find(qn('c:spPr'))
    assert spPr is None
