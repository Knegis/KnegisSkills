"""End-to-end: build_deck.py on real reference HTML through the full pipeline.

Marked @pytest.mark.browser - these need Playwright + a real Edge/Chrome to run
extract_geometry.py (see scripts/extract_geometry.py). Run explicitly with:
    python -m pytest tests -m browser
(the default `-q` run in this suite's other files excludes this marker).
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import fontcheck  # noqa: E402

from conftest import SKILL_ROOT, run_py  # noqa: E402

REFERENCE = SKILL_ROOT / "reference"
T1 = REFERENCE / "T1_statement_hero_stat.html"
T2 = REFERENCE / "T2_trend_stackedbar_driver_rail.html"
T3 = REFERENCE / "T3_evidence_source_rows_commentary.html"

pytestmark = pytest.mark.browser


def test_build_deck_two_reference_slides(tmp_path):
    out = tmp_path / "deck.pptx"
    r = run_py([str(SKILL_ROOT / "scripts" / "build_deck.py"), str(out), str(T1), str(T2)],
               cwd=SKILL_ROOT, timeout=180)
    assert r.returncode == 0, r.stdout + r.stderr
    assert out.exists()

    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 2

    for slide in prs.slides:
        assert slide.slide_layout.name == "Blank"

    has_chart = any(shape.has_chart for slide in prs.slides for shape in slide.shapes)
    assert has_chart, "expected T2's .ppt-chart to become a real, editable python-pptx chart"

    target, seen, offenders = fontcheck.scan(out)
    assert offenders == [], f"font offenders: {offenders}"


def test_typecheck_passes_on_t3_reference_slide():
    r = run_py([str(SKILL_ROOT / "scripts" / "typecheck.py"), str(T3), "--quiet"],
               cwd=SKILL_ROOT, timeout=120)
    assert r.returncode == 0, r.stdout + r.stderr
