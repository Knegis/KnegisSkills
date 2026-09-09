"""Generate the neutral 16:9 template the skill ships with.

Usage: python make_template.py [out.pptx]       (default: ../template/house_template.pptx)

Starts from python-pptx's built-in default presentation (11 standard layouts: Title Slide,
Title and Content, Section Header, Two Content, Comparison, Title Only, Blank, ...), switches
it to 16:9 (13.333 x 7.5 in), stretches every layout placeholder horizontally so the 4:3
positions still make sense, sets the theme fonts to the house font, and strips authoring
metadata. No logo, no footer text, no photography - bring your own template for those, or
set wordmark_text in house.json and the HTML boilerplate draws it.
"""
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as C  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

W43, W169 = Inches(10), Inches(13.333)


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else C.SKILL_ROOT / "template" / "house_template.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W169, Inches(7.5)
    k = W169 / W43
    for layout in prs.slide_layouts:
        for shp in layout.placeholders:
            if shp.left is not None and shp.width is not None:
                shp.left = int(shp.left * k)
                shp.width = int(shp.width * k)
    for shp in prs.slide_master.placeholders:
        if shp.left is not None and shp.width is not None:
            shp.left = int(shp.left * k)
            shp.width = int(shp.width * k)
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.title = ""
    prs.save(str(out))

    # theme fonts + any hardcoded typefaces -> house font; theme name neutral
    tmp = Path(tempfile.mkdtemp())
    with zipfile.ZipFile(out) as z:
        z.extractall(tmp)
    for p in tmp.rglob("*.xml"):
        s = p.read_text(encoding="utf-8")
        s2 = re.sub(r'typeface="(Calibri|Calibri Light|Arial|Helvetica)"', f'typeface="{C.FONT_NAME}"', s)
        s2 = s2.replace('name="Office Theme"', 'name="House Theme"')
        if s2 != s:
            p.write_text(s2, encoding="utf-8")
    import office_io
    office_io.pack(tmp, out)
    shutil.rmtree(tmp, ignore_errors=True)

    prs = Presentation(str(out))
    print(f"template written: {out}")
    print(f"  {prs.slide_width / 914400:.3f} x {prs.slide_height / 914400:.2f} in, "
          f"{len(prs.slide_layouts)} layouts, font {C.FONT_NAME}")
    for i, l in enumerate(prs.slide_layouts, 1):
        print(f"  slideLayout{i}.xml  {l.name}")


if __name__ == "__main__":
    main()
